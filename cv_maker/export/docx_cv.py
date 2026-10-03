from pathlib import Path

from docx import Document
from docx.shared import Pt

from cv_maker.honesty import CvDocument


class ExportError(Exception):
    pass


def verify_docx(path: Path, must_contain: str = "") -> None:
    """Open what was just written, before it replaces anything: a file Word couldn't open, or one missing its text,
    is an error, never a download."""
    try:
        text = "\n".join(p.text for p in Document(str(path)).paragraphs)
    except Exception as exc:
        raise ExportError(f"the file just written couldn't be opened again ({type(exc).__name__})") from exc
    if not text.strip() or (must_contain and must_contain not in text):
        raise ExportError("the file just written is missing its text")


def _set_calibri(run, size: float = 11, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.bold = bold or None
    run.italic = italic or None


def _para(document, text: str, *, size: float = 11, bold: bool = False, italic: bool = False, style: str | None = None):
    p = document.add_paragraph(style=style)
    _set_calibri(p.add_run(text), size=size, bold=bold, italic=italic)
    return p


def _dates(start: str, end: str | None, current: bool) -> str:
    finish = "Present" if current else (end or "")
    return " – ".join(x for x in (start, finish) if x)


def write_cv_docx(doc: CvDocument, path: Path) -> None:
    """ATS-safe single-column DOCX: Word heading styles, no tables, no header/footer text."""
    tmp_path = path.with_suffix(".tmp.docx")
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    # Some ATS read the document's properties: your name as the author, the role as the title.
    document.core_properties.author = doc.name
    document.core_properties.title = f"{doc.name} - {doc.headline}".strip(" -") if doc.headline else f"{doc.name} CV"
    document.core_properties.comments = ""
    try:
        _para(document, doc.name, size=20, bold=True)
        if doc.headline:
            _para(document, doc.headline, size=12)
        if doc.contact:
            _para(document, " | ".join(doc.contact), size=10)

        if doc.summary:
            document.add_heading("Summary", level=1)
            _para(document, doc.summary)

        # Skills near the top: the posting's keywords are the first thing an ATS and a recruiter look for.
        document.add_heading("Skills", level=1)
        _para(document, ", ".join(doc.skills))

        document.add_heading("Experience", level=1)
        for exp in doc.experiences:
            _para(document, " — ".join(x for x in (exp.title, exp.company) if x), bold=True)
            meta = " | ".join(x for x in (exp.location, _dates(exp.start, exp.end, exp.current)) if x)
            if meta:
                _para(document, meta, size=10, italic=True)
            for bullet in exp.bullets:
                _para(document, bullet, style="List Bullet")

        document.add_heading("Education", level=1)
        for edu in doc.education:
            title = ", ".join(x for x in (edu.credential, edu.field) if x)
            line = " — ".join(x for x in (title, edu.school) if x)
            _para(document, f"{line} ({edu.dates})" if edu.dates else line)

        for key, heading in (("certifications", "Certifications"), ("projects", "Projects"), ("languages", "Languages")):
            items = doc.extras.get(key) or []
            if items:
                document.add_heading(heading, level=1)
                for item in items:
                    _para(document, item, style="List Bullet")

        document.save(tmp_path)
        verify_docx(tmp_path, doc.name)
        tmp_path.replace(path)
    except Exception as exc:
        if tmp_path.exists():
            tmp_path.unlink()
        if isinstance(exc, ExportError):
            raise
        raise ExportError(str(exc)) from exc
