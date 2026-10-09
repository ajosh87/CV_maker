from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt

from cv_maker.export.docx_cv import ExportError, verify_docx


def write_letter_docx(body: str, path: Path, name: str = "", contact: list[str] | None = None) -> None:
    tmp_path = path.with_suffix(".tmp.docx")
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    def para(text: str, size: float = 11, bold: bool = False) -> None:
        run = document.add_paragraph().add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.bold = bold or None

    try:
        if name:
            para(name, size=14, bold=True)
        if contact:
            para(" | ".join(contact), size=10)
        if name or contact:
            para(date.today().strftime("%d %B %Y"))
        for block in [b.strip() for b in body.replace("\r\n", "\n").split("\n\n")]:
            if block:
                para(block)
        document.save(tmp_path)
        verify_docx(tmp_path, name)
        tmp_path.replace(path)
    except Exception as exc:
        if tmp_path.exists():
            tmp_path.unlink()
        if isinstance(exc, ExportError):
            raise
        raise ExportError(str(exc)) from exc
