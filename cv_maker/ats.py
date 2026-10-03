"""The ATS check: how an applicant tracking system reads this CV for this job.

Tracking systems rank CVs mostly by the posting's own keywords: whether they appear, in the posting's wording, and
whether they appear in context (a role's bullets) rather than only in a skills list. They also need standard sections,
a title that matches, contact details they can parse, and dates on every role. Everything here is counted from the
CV's text, so each point can be traced; no LLM grades it.

The same keyword plan steers the writing (which of the posting's words you can honestly use, and where your profile
supports them), and `optimise` makes sure every keyword you have appears at least once in the posting's wording. A
keyword you don't have is never added: it stays an honest gap, shown as one.
"""
import re

from cv_maker import skills as sk

WEIGHTS = {"must": 45, "nice": 10, "context": 10, "title": 8, "sections": 7, "contact": 5, "dates": 5, "numbers": 5,
           "length": 5}
_TITLE_NOISE = re.compile(r"\s*(?:[-–|,]\s*(?:remote|hybrid|on-?site|full[- ]time|part[- ]time|contract|permanent)\b.*|"
                          r"\((?:m/f/d|f/m/d|m/w/d|w/m/d|all genders|remote|hybrid)\)|\s+-\s+[A-Z][\w ]+,\s*[A-Z][\w ]+$)", re.I)
_STOP = {"and", "or", "the", "of", "for", "with", "in", "a", "an", "to", "senior", "junior", "lead", "principal", "staff",
         "i", "ii", "iii", "sr", "jr"}


def clean_title(title: str, company: str = "") -> str:
    """'AI & Cloud Solutions Engineer - Remote (m/f/d)' -> 'AI & Cloud Solutions Engineer'; 'Engineer at Acme' -> 'Engineer'."""
    title = " ".join(_TITLE_NOISE.sub("", title or "").split()).strip(" -–|,")
    if company:
        title = re.sub(rf"\s+(?:at|@|-|–|\|)\s+{re.escape(company)}\s*$", "", title, flags=re.I).strip()
    return title


def plan(requirements: dict, gaps: list[dict]) -> list[dict]:
    """The posting's keywords in its own words: must-haves first, with whether you have them and why."""
    by_term = {g["term"].casefold(): g for g in gaps or []}
    out = []
    for key, required in (("must_have", True), ("nice_to_have", False)):
        for item in (requirements or {}).get(key) or []:
            term = item.get("term", "")
            gap = by_term.get(term.casefold(), {})
            have = bool(gap.get("level") in sk.LEVELS or (gap.get("same") and int(gap.get("score") or 0) >= sk.PARTIAL))
            support = [f["text"] for f in gap.get("found") or [] if f.get("where") not in ("Your answer",)][:2]
            out.append({"term": term, "required": required, "have": have, "level": gap.get("level", ""),
                        "score": int(gap.get("score") or 0), "support": support})
    return out


def prompt_lines(keyword_plan: list[dict]) -> str:
    """The plan, as the writing prompt reads it."""
    have = [k for k in keyword_plan if k["have"]]
    lacking = [k["term"] for k in keyword_plan if not k["have"]]
    lines = []
    if have:
        lines.append("Keywords from the posting that the candidate HAS. Use the posting's exact wording for each at least once, "
                     "in a bullet or the summary where a fact below supports it (applicant tracking systems match exact words):")
        for k in have:
            level = f", {k['level']} level" if k["level"] else ""
            support = "; ".join(f"“{s[:90]}”" for s in k["support"]) or "confirmed by the candidate"
            lines.append(f"- {k['term']} ({'must-have' if k['required'] else 'nice-to-have'}{level}) — supported by: {support}")
    if lacking:
        lines.append(f"Keywords the candidate does NOT have (never claim or imply them): {', '.join(lacking)}")
    return "\n".join(lines)


def _expanded(term: str) -> str:
    """'RAG' -> 'Retrieval-Augmented Generation (RAG)': both forms, for whichever a recruiter searches."""
    full = sk.ACRONYMS.get(sk.norm(term))
    return f"{full} ({term})" if full and term.isupper() else term


def optimise(doc, keyword_plan: list[dict], job_title: str = "", company: str = ""):
    """Put every keyword you have into the skills list in the posting's wording (first, in the posting's order),
    with acronyms written out once, and title the CV for the job. Never adds a keyword you don't have."""
    skills = list(doc.skills)
    ranked: list[str] = []
    for k in keyword_plan:
        if not k["have"]:
            continue
        term = k["term"]
        if any(sk.variants(term) & sk.variants(s) or sk.norm(s) == sk.norm(_expanded(term)) for s in ranked):
            continue
        twin = next((s for s in skills if sk.variants(term) & sk.variants(s)), None)
        if twin is not None:
            skills.remove(twin)  # the same skill under another name: the posting's wording replaces yours
        ranked.append(_expanded(term))
    rest = [s for s in skills if not any(sk.variants(s) & sk.variants(r) or sk.found_in(s, r) for r in ranked)]
    doc.skills = (ranked + rest)[:40]
    title = clean_title(job_title, company)
    if title:
        doc.headline = title
    return doc


# ---- reading the CV the way an ATS does -----------------------------------------------------------------------

def _parts(cv: dict) -> list[tuple[str, str]]:
    """(where, text) for every part of the CV an ATS reads."""
    parts = [("headline", cv.get("headline") or ""), ("summary", cv.get("summary") or ""),
             ("skills", ", ".join(cv.get("skills") or []))]
    for e in cv.get("experiences") or []:
        role = e.get("company") or e.get("title") or "a role"
        parts.append((f"{role} (title)", e.get("title") or ""))
        parts += [(f"{role} (bullet {i})", b) for i, b in enumerate(e.get("bullets") or [], 1)]
    for key in ("certifications", "projects", "languages"):
        parts += [(key, str(x)) for x in (cv.get("extras") or {}).get(key) or []]
    for e in cv.get("education") or []:
        parts.append(("education", f"{e.get('credential', '')} {e.get('field', '')}"))
    return parts


def _count(term: str, text: str) -> int:
    return max((len(sk._pattern(v).findall(text)) for v in sk.variants(term)), default=0)


def check(cv: dict, keyword_plan: list[dict], job_title: str = "", company: str = "") -> dict:
    """{"score", "grade", "checks": [...], "keywords": [...]}: every number from the CV's own text."""
    parts = _parts(cv)
    whole = "\n".join(text for _, text in parts)
    keywords = []
    for k in keyword_plan:
        places = [where for where, text in parts if text and sk.found_in(k["term"], text)]
        exact = any(re.search(rf"(?<![\w+#]){re.escape(k['term'])}(?![\w+#])", text, re.I) for _, text in parts)
        keywords.append({**k, "places": places, "count": _count(k["term"], whole), "in_cv": bool(places), "exact": exact,
                         "in_context": any(p.startswith("summary") or "(bullet" in p for p in places)})
    must = [k for k in keywords if k["required"]]
    nice = [k for k in keywords if not k["required"]]
    checks = []

    def add(key, label, ratio, detail, fix=""):
        ratio = max(0.0, min(1.0, ratio))
        checks.append({"key": key, "label": label, "points": round(WEIGHTS[key] * ratio, 1), "max": WEIGHTS[key],
                       "ok": ratio >= 0.999, "detail": detail, "fix": fix if ratio < 0.999 else ""})

    found_must = [k for k in must if k["in_cv"]]
    missing_have = [k["term"] for k in must + nice if k["have"] and not k["in_cv"]]
    lacking = [k["term"] for k in must if not k["have"] and not k["in_cv"]]
    add("must", "Must-have keywords", len(found_must) / len(must) if must else 1,
        f"{len(found_must)} of {len(must)} in your CV" if must else "The posting lists no must-haves",
        (f"Missing though you have them: {', '.join(missing_have)}. Write a new version." if missing_have else "") +
        (f" Not in your profile: {', '.join(lacking)}. If you do have any, answer its question on a new version."
         if lacking else ""))
    found_nice = [k for k in nice if k["in_cv"]]
    add("nice", "Nice-to-have keywords", len(found_nice) / len(nice) if nice else 1,
        f"{len(found_nice)} of {len(nice)} in your CV" if nice else "The posting lists none")
    in_context = [k for k in found_must if k["in_context"]]
    add("context", "Keywords used in context", len(in_context) / len(found_must) if found_must else 1,
        f"{len(in_context)} of {len(found_must)} found must-haves appear in a bullet or the summary, not only the skills list",
        "Keywords only in the skills list count for less. Add the facts that show them to your roles on the Profile page.")
    title = clean_title(job_title, company)
    words = [w for w in re.findall(r"[\w+#.]+", title.lower()) if w not in _STOP]
    head = f"{cv.get('headline') or ''} {cv.get('summary') or ''}".lower()
    title_ratio = 1.0 if not words or title.lower() in head else sum(w in head for w in words) / len(words)
    add("title", "Job title", title_ratio if title_ratio >= 0.6 else title_ratio / 2,
        f"“{title}” {'appears' if title_ratio >= 0.999 else 'partly appears' if title_ratio else 'does not appear'} at the top"
        if title else "The posting has no title", "Use the posting's job title as your CV's headline.")
    sections = [("Summary", bool(cv.get("summary"))), ("Experience", bool(cv.get("experiences"))),
                ("Skills", bool(cv.get("skills"))), ("Education", bool(cv.get("education")))]
    have_sections = [name for name, ok in sections if ok]
    add("sections", "Standard sections", len(have_sections) / len(sections),
        f"{', '.join(have_sections) or 'None'}" + (f"; missing {', '.join(n for n, ok in sections if not ok)}"
                                                    if len(have_sections) < len(sections) else ""),
        "Add your education on the Profile page." if not cv.get("education") else "")
    contact = " ".join(cv.get("contact") or [])
    email = bool(re.search(r"[^@\s]+@[^@\s]+\.\w+", contact))
    phone = len(re.sub(r"\D", "", contact)) >= 7
    add("contact", "Contact details", (email + phone) / 2,
        "Email and phone at the top" if email and phone else f"Missing: {', '.join(x for x, ok in (('email', email), ('phone', phone)) if not ok)}",
        "Add them on the Profile page: an ATS that can't find them may drop the CV.")
    roles = cv.get("experiences") or []
    dated = [e for e in roles if e.get("start")]
    add("dates", "Dates on every role", len(dated) / len(roles) if roles else 1,
        f"{len(dated)} of {len(roles)} roles have dates", "Add start (and end) dates to every role on the Profile page.")
    bullets = [b for e in roles for b in e.get("bullets") or []]
    numbered = [b for b in bullets if re.search(r"\d", b)]
    ratio = len(numbered) / len(bullets) if bullets else 0
    add("numbers", "Measurable results", min(1.0, ratio / 0.3), f"{len(numbered)} of {len(bullets)} bullets have numbers",
        "Recruiters and ATS rankings favour results with numbers. Add the ones you have to your bullets on the Profile page.")
    count = len(re.findall(r"\w+", whole))
    length = 1.0 if 300 <= count <= 1100 else 0.6 if 200 <= count <= 1400 else 0.2
    add("length", "Length", length, f"About {count} words ({'one to two pages' if length == 1 else 'short' if count < 300 else 'long'})",
        "Aim for 300 to 1,100 words: one to two pages.")
    score = round(sum(c["points"] for c in checks))
    stuffed = [k["term"] for k in keywords if k["count"] > 8]
    return {"score": score, "grade": "Excellent" if score >= 85 else "Good" if score >= 70 else "Fair" if score >= 50 else "Weak",
            "checks": checks, "keywords": keywords, "missing_have": missing_have, "lacking": lacking, "stuffed": stuffed,
            "title": title}


# ---- the file itself, read back the way an ATS parser reads it -------------------------------------------------

_HEADINGS = (("Summary", "summary"), ("Skills", "skills"), ("Experience", "experiences"), ("Education", "education"))


def read_back(path: str, cv: dict, result: dict) -> dict:
    """Open the DOCX the employer gets and read it as a parser would: is everything the check counted really in the
    file, where a parser looks for it? {"ok": bool, "items": [{"label", "ok", "detail"}]}."""
    from docx import Document

    items = []

    def add(label: str, ok, detail: str) -> None:
        items.append({"label": label, "ok": bool(ok), "detail": detail})

    try:
        document = Document(path)
    except Exception as exc:  # a file that won't open is the one thing worse than a low score
        return {"ok": False, "items": [{"label": "Opens as a Word document", "ok": False,
                                        "detail": f"it couldn't be opened ({type(exc).__name__})"}]}
    paragraphs = [((p.style.name if p.style is not None else "") or "", p.text) for p in document.paragraphs]
    text = "\n".join(t for _, t in paragraphs)
    add("Reads as plain text", text.strip(), f"{len(re.findall(r'[A-Za-z0-9]+', text)):,} words, in reading order")
    headings = {t.strip().casefold() for style, t in paragraphs if style.startswith("Heading")}
    wanted = [name for name, key in _HEADINGS if cv.get(key)]
    missing = [name for name in wanted if name.casefold() not in headings]
    add("Standard headings, marked as headings", not missing,
        f"{', '.join(wanted)}" if not missing else f"not marked as headings: {', '.join(missing)}")
    contact = " ".join(cv.get("contact") or [])
    email = re.search(r"[^@\s|]+@[^@\s|]+\.\w+", contact)
    phone = re.search(r"\+?\d[\d\s().-]{6,}\d", contact)
    in_body = [label for label, found in (("email", email), ("phone", phone)) if found and found.group(0) in text]
    lost = [label for label, found in (("email", email), ("phone", phone)) if found and found.group(0) not in text]
    add("Contact details in the body text", not lost,
        (f"{' and '.join(in_body)} found in the text, not in a header or footer" if in_body else "none to check") if not lost
        else f"not found in the text: {', '.join(lost)}")
    counted = [k["term"] for k in result.get("keywords") or [] if k.get("in_cv")]
    absent = [term for term in counted if not sk.found_in(term, text)]
    add("Every keyword counted is in the file", not absent,
        f"all {len(counted)} found in the file's text" if not absent else f"missing from the file: {', '.join(absent)}")
    xml = document.element.xml
    extras = [what for what, found in (("tables", document.tables), ("text boxes", "txbxContent" in xml),
                                        ("images", "<w:drawing" in xml or "v:imagedata" in xml),
                                        ("columns", re.search(r'<w:cols\b[^>]*w:num="([2-9])"', xml))) if found]
    header_text = " ".join(p.text for s in document.sections for part in (s.header, s.footer) for p in part.paragraphs).strip()
    if header_text:
        extras.append("text in the header or footer")
    add("One column, nothing a parser skips", not extras,
        "no tables, text boxes, images, columns or header text" if not extras else f"has {', '.join(extras)}")
    return {"ok": all(i["ok"] for i in items), "items": items}
