"""The apply step in the app: reached from a job's page, optional, observable, and cleaned up with your data."""
import json
import time
from io import BytesIO

import pytest
from docx import Document

from cv_maker.app import create_app
from cv_maker.apply import accounts
from cv_maker.jobs.fetch import FetchResult

PROFILE = {"name": "Ada Lovelace", "email": "ada@example.com", "phone": "+44 20 7946 0958", "location": "London", "links": {},
           "experiences": [{"company": "Northwind", "title": "Engineer", "start": "2021", "end": None, "current": True,
                            "bullets": ["Built Python APIs"]}], "education": [], "skills": ["Python"]}


class FakeModel:
    last_usage = {"in": 1200, "out": 300}  # like a provider that reports usage

    def complete(self, messages, *, json_mode=False):
        prompt = messages[-1]["content"]
        if prompt.startswith("Extract a structured profile"):
            return json.dumps(PROFILE)
        if prompt.startswith("Read this job description"):
            return json.dumps({"job_title": "Platform Engineer", "company": "Acme", "must_have": [{"term": "Python", "original": "Python"}]})
        if prompt.startswith("You are tailoring"):
            return json.dumps({"summary": "Python engineer.", "skills": ["Python"],
                               "experiences": [{"company": "Northwind", "title": "Engineer", "bullets": ["Built Python APIs"]}]})
        return json.dumps({"next": {"do": "wait"}})


def _app(tmp_path, url="https://jobs.example.com/1"):
    fetch = lambda u: FetchResult(ok=True, text="Need Python", reason="", title="Platform Engineer", company="Acme")  # noqa: E731
    return create_app(data_dir=tmp_path, model_factory=FakeModel, fetcher=fetch, sync_jobs=True)


def _ready_job(app, client, url="https://jobs.example.com/1"):
    buf = BytesIO()
    doc = Document()
    doc.add_paragraph("Ada Lovelace")
    doc.add_paragraph("Engineer at Northwind since 2021. Built Python APIs used by three teams and ran on-call for the payments platform.")
    doc.save(buf)
    res = client.post("/profile/upload", data={"cv": (BytesIO(buf.getvalue()), "cv.docx")}, content_type="multipart/form-data")
    client.post(res.headers["Location"], data={"action": "send"})
    client.post("/jobs", data={"urls": url})
    run = app.store.list_runs()[0]
    client.post(f"/runs/{run.id}/questions", data={})
    return app.store.get_run(run.id)


def test_applying_is_an_optional_next_step_on_the_job_page(tmp_path):
    app = _app(tmp_path)
    client = app.test_client()
    run = _ready_job(app, client)
    page = client.get(f"/runs/{run.id}").data.decode()
    assert 'id="apply"' in page and "Stop here and apply yourself" in page and f"/runs/{run.id}/apply" in page
    start = client.get(f"/runs/{run.id}/apply").data.decode()
    assert "Apply with the assistant" in start and 'value="https://jobs.example.com/1"' in start
    assert "never submits by itself" in start and "never screenshots" in start
    assert client.post(f"/runs/{run.id}/apply", data={"start_url": "javascript:alert(1)"}).status_code == 302
    assert app.store.list_applications() == []  # refused


def test_no_cv_means_nothing_to_apply_with(tmp_path):
    app = _app(tmp_path)
    run = app.store.create_run(job_url="https://x")
    res = app.test_client().get(f"/runs/{run.id}/apply")
    assert res.status_code == 302 and res.headers["Location"].endswith(f"/runs/{run.id}")


def test_marking_an_application_as_submitted_shows_on_the_job_and_the_list(tmp_path):
    app = _app(tmp_path)
    client = app.test_client()
    run = _ready_job(app, client)
    application = app.store.create_application(run, 1, "https://jobs.example.com/1")
    application.status = "stopped"
    app.store.update_application(application)
    body = client.get(f"/api/applications/{application.id}").get_json()
    assert body["status"] == "stopped" and body["live"] is False
    assert client.post(f"/api/applications/{application.id}/control", json={"action": "resume"}).status_code == 409
    assert client.post(f"/api/applications/{application.id}/control", json={"action": "mark_submitted"}).get_json()["ok"]
    assert "Applied" in client.get(f"/runs/{run.id}").data.decode()
    assert "Applied" in client.get("/jobs").data.decode()
    page = client.get(f"/applications/{application.id}").data.decode()
    assert 'id="ap-live"' in page and "Filled so far" in page


def test_application_details_are_saved_and_contact_details_are_marked_private(tmp_path):
    app = _app(tmp_path)
    client = app.test_client()
    _ready_job(app, client)
    page = client.get("/settings/apply").data.decode()
    assert 'value="Ada"' in page and 'value="Lovelace"' in page  # started from the profile
    assert page.count("hidden from the LLM") >= 8
    client.post("/settings/apply", data={"first_name": "Ada", "last_name": "Lovelace", "email": "ada@example.com",
                                         "salary": "£70,000", "custom": "Years of experience with Python = 6\nnot a pair"})
    page = client.get("/settings/apply").data.decode()
    assert "£70,000" in page and "Years of experience with Python = 6" in page
    assert "not a pair” has no “question = answer”" in page  # said, not silently dropped
    assert "not a pair" not in page.split('id="custom"')[1].split("</textarea>")[0]  # and not saved


def test_forgetting_an_account_and_deleting_everything(tmp_path):
    app = _app(tmp_path)
    client = app.test_client()
    _ready_job(app, client)
    accounts.remember_account(app.store, "myworkdayjobs.com", "ada@example.com", in_keychain=False)
    assert "myworkdayjobs.com" in client.get("/settings/apply").data.decode()
    client.post("/settings/apply/forget", data={"site": "myworkdayjobs.com", "email": "ada@example.com"})
    assert accounts.list_accounts(app.store) == []

    accounts.remember_account(app.store, "acme.com", "ada@example.com", in_keychain=False)
    (tmp_path / "browser").mkdir()
    (tmp_path / "browser" / "Cookies").write_bytes(b"x")
    run = app.store.list_runs()[0]
    app.store.create_application(run, 1, "https://x")
    client.post("/settings/delete-all", data={"confirm": "DELETE"})
    assert not (tmp_path / "browser").exists()
    assert app.store.list_applications() == [] and accounts.list_accounts(app.store) == []
    assert app.store.get_preference("apply_details") is None


def test_passwords_go_to_a_password_manager_only_when_you_confirm(tmp_path, monkeypatch):
    app = _app(tmp_path)
    client = app.test_client()
    vault = {("myworkdayjobs.com", "ada@example.com"): "S3cret=pw+1"}
    monkeypatch.setattr(accounts, "can_store", lambda: True)
    monkeypatch.setattr(accounts, "get_password", lambda site, email: vault.get((site, email)))
    accounts.remember_account(app.store, "myworkdayjobs.com", "ada@example.com", in_keychain=True)
    accounts.remember_account(app.store, "acme.com", "ada@example.com", in_keychain=False)  # its password was shown once
    page = client.get("/settings/apply").data.decode()
    assert "Download for a password manager" in page and "with the 1 saved password" in page and "S3cret" not in page

    refused = client.post("/settings/apply/passwords.csv", data={}, follow_redirects=True)
    assert "Tick the box first" in refused.data.decode()
    res = client.post("/settings/apply/passwords.csv", data={"understood": "1"})
    assert res.headers["Cache-Control"] == "no-store" and "attachment" in res.headers["Content-Disposition"]
    lines = res.data.decode().splitlines()
    assert lines[0] == "name,url,username,password,note"
    assert lines[1].startswith("myworkdayjobs.com,https://myworkdayjobs.com/,ada@example.com,S3cret=pw+1,") and len(lines) == 2
    events = client.get("/api/events").get_json()["events"]
    assert any("exported 1 job-site password" in e["text"] for e in events) and not any("S3cret" in e["text"] for e in events)
    assert "S3cret" not in (tmp_path / "logs" / "cv-tailor.log").read_text(encoding="utf-8")


def test_nerdbar_shows_what_ran_how_long_and_the_tokens(tmp_path):
    app = _app(tmp_path)
    client = app.test_client()
    before = client.get("/api/events").get_json()["events"]
    after_seq = before[-1]["seq"] if before else 0
    _ready_job(app, client)
    data = client.get(f"/api/events?after={after_seq}").get_json()
    texts = [e["text"] for e in data["events"]]
    assert any(t.startswith("▶ Reading cv.docx") for t in texts) and any(t.startswith("■ Writing the CV") for t in texts)
    assert any(t.startswith("→ Reading your CV") and "hid" in t for t in texts)
    assert any(t.startswith("← 1,200 in · 300 out") for t in texts)  # usage as the provider reported it
    assert all("ada@example.com" not in t and "Lovelace" not in t for t in texts)
    tags = {e["tag"] for e in data["events"]}
    assert any(tag.startswith("job ") for tag in tags) and any(tag.startswith("cv ") for tag in tags)
    assert data["stats"]["session"]["tokens_in"] >= 3600 and data["stats"]["session"]["llm_calls"] >= 3
    page = client.get("/jobs").data.decode()
    assert 'id="nb-panel"' in page and 'id="nb-toggle"' in page and "nerdbar.js" in page
    sent = client.get("/settings/sent").data.decode()
    assert "1,200 tokens in, 300 out" in sent


# ---- through the app, on the mock site ----

def _browser_available() -> bool:
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as pw:
            pw.chromium.launch(headless=True).close()
        return True
    except Exception:
        return False


@pytest.mark.skipif(not _browser_available(), reason="Playwright with Chromium is not installed")
def test_start_watch_and_stop_from_the_app(tmp_path):
    from fixtures.mock_ats import Running

    site = Running()
    app = _app(tmp_path)
    client = app.test_client()
    try:
        run = _ready_job(app, client)
        res = client.post(f"/runs/{run.id}/apply", data={"start_url": f"{site.url}/job/1", "version": "1"})
        app_id = res.headers["Location"].rsplit("/", 1)[1]
        end = time.monotonic() + 60
        while time.monotonic() < end:
            state = client.get(f"/api/applications/{app_id}").get_json()
            if (state["waiting"] or {}).get("kind") == "account":
                break
            time.sleep(0.3)
        assert state["waiting"]["kind"] == "account" and state["live"] and state["mode"] == "paused"
        frame = client.get(f"/api/applications/{app_id}/frame")
        assert frame.status_code == 200 and frame.mimetype == "image/jpeg" and frame.data[:2] == b"\xff\xd8"
        assert client.post(f"/runs/{run.id}/apply", data={"start_url": f"{site.url}/job/1"}).status_code == 302
        assert len(app.store.list_applications()) == 1  # one assistant at a time
        client.post(f"/api/applications/{app_id}/control", json={"action": "stop"})
        app.apply_sessions.stop_all()
        assert app.store.get_application(app_id).status == "stopped"
        assert client.get(f"/api/applications/{app_id}/frame").status_code == 200  # the last view stays
    finally:
        app.apply_sessions.stop_all()
        site.close()
