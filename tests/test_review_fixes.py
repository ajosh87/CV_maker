"""Regression tests for the pre-push code review findings."""
import io
import json
import warnings
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document

from cv_maker.app import create_app
from cv_maker.honesty import CvDocument, FactCheck, filter_cv_document
from cv_maker.jobs.fetch import FetchResult
from cv_maker.models import Experience, Profile, Skill
from cv_maker.pipeline.graph import Pipeline, apply_answers_node
from cv_maker.settings import effective_settings
from cv_maker.store import Store
from cv_maker.tasks import Runner, Tasks

PROFILE_JSON = {"name": "Ada", "skills": ["Python"], "experiences": [{"company": "Northwind", "title": "Engineer", "bullets": ["Built APIs"]}]}


def _profile(**kw):
    base = dict(name="Ada", email="", phone="", location="", links={}, target_role=None,
                experiences=[Experience("Northwind", "Engineer", "", "2021", None, True, ["Built APIs for 3 teams"], True)],
                education=[], skills=[Skill("Python", "cv")], extras={})
    base.update(kw)
    return Profile(**base)


# ---- deleting while background work runs ----

class DeletesDuringCall:
    """Simulates the user clicking "Delete all my data" while the LLM is answering."""

    def __init__(self, store):
        self.store = store

    def complete(self, messages, *, json_mode=False):
        self.store.delete_everything()
        if messages[-1]["content"].startswith("Extract a structured profile"):
            return json.dumps(PROFILE_JSON)
        return json.dumps({"summary": "Python developer.", "experiences": [], "skills": ["Python"]})


def test_delete_all_during_cv_parsing_does_not_bring_the_profile_back(tmp_path):
    store = Store(tmp_path / "db.sqlite")
    model = DeletesDuringCall(store)
    tasks = Tasks(store, Pipeline(store, lambda: model, tmp_path), lambda: model, None, Runner(sync=True))
    tasks.queue_parse(store.create_upload("cv.docx", "x", "Ada, Python"))
    assert store.get_profile() is None


def test_delete_during_writing_leaves_no_files(tmp_path):
    store = Store(tmp_path / "db.sqlite")
    store.save_profile(_profile())
    run = store.create_run(job_url="https://x")
    run.jd_text, run.requirements = "Python", {"must_have": []}
    store.update_run(run)
    model = DeletesDuringCall(store)
    result = Pipeline(store, lambda: model, tmp_path).generate(run.id)
    assert result is None
    assert list(tmp_path.glob("*.docx")) == []


# ---- free-text fact check ----

def test_summary_with_invented_numbers_or_employers_is_flagged():
    check = FactCheck()
    doc = CvDocument("", "Engineer with 15 years at Google leading 200 people. Python developer at Northwind.", [], [], ["Python"])
    out = filter_cv_document(doc, _profile(), [], check)
    assert out.summary == "Python developer at Northwind."
    assert "15, 200" in check.items[0]["reason"]


def test_summary_names_must_come_from_profile_or_job():
    check = FactCheck()
    doc = CvDocument("", "Python engineer ready to join AstraZeneca. Former Google engineer.", [], [], ["Python"])
    out = filter_cv_document(doc, _profile(), [], check, context_text="AI Engineer AstraZeneca")
    assert out.summary == "Python engineer ready to join AstraZeneca."
    assert check.items[0]["reason"].startswith("names Google")


def test_letter_sentences_are_checked_too(tmp_path):
    class Model:
        def complete(self, messages, *, json_mode=False):
            p = messages[-1]["content"]
            if p.startswith("Write a cover letter"):
                return json.dumps({"letter": "Dear team,\n\nI built APIs for 3 teams. I grew revenue by 40% at Google.\n\nAda"})
            return json.dumps({"summary": "Python developer.", "experiences": [], "skills": ["Python"]})

    store = Store(tmp_path / "db.sqlite")
    store.save_profile(_profile())
    run = store.create_run(job_url="https://x")
    run.jd_text, run.requirements, run.generate_letter = "Python", {"must_have": []}, True
    store.update_run(run)
    done = Pipeline(store, lambda: Model(), tmp_path).generate(run.id)
    text = "\n".join(p.text for p in Document(done.letter_path).paragraphs)
    assert "I built APIs for 3 teams." in text and "Google" not in text


# ---- concurrency & storage ----

def test_store_closes_its_connections(tmp_path):
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        store = Store(tmp_path / "db.sqlite")
        store.create_run(job_url="x")
        store.list_runs()
        import gc
        gc.collect()
    assert not [w for w in caught if issubclass(w.category, ResourceWarning)]


def test_answers_are_applied_to_the_current_profile(tmp_path):
    """Another job confirmed 'Go' after this job loaded its copy of the profile; it must survive."""
    store = Store(tmp_path / "db.sqlite")
    stale = _profile()
    store.save_profile(_profile(skills=[Skill("Python", "cv"), Skill("Go", "clarification")]))
    state = {
        "run_id": "r", "profile": json.loads(json.dumps(stale.__dict__, default=lambda o: o.__dict__)),
        "answers": [{"term": "Rust", "skipped": False, "text": "yes", "add_role": None}],
        "questions": [{"term": "Rust", "prompt": "Rust?", "kind": "yesno", "required": True}],
        "requirements": {"must_have": [{"term": "Rust", "original": "Rust"}]},
    }
    from cv_maker.pipeline.graph import PipelineContext

    apply_answers_node(state, PipelineContext(store, lambda: None, tmp_path))
    assert {s.name for s in store.get_profile().skills} == {"Python", "Go", "Rust"}


def test_version_numbers_are_never_reused(tmp_path):
    store = Store(tmp_path / "db.sqlite")
    run = store.create_run(job_url="x")
    store.add_document(run, "cv", 1, "a")
    doc2 = store.add_document(run, "cv", 2, "b")
    run.last_version = 2
    store.update_run(run)
    store.delete_document(doc2.id)
    assert store.next_version(run.id) == 3


# ---- app-level fixes ----

class FakeModel:
    def complete(self, messages, *, json_mode=False):
        p = messages[-1]["content"]
        if p.startswith("Extract a structured profile"):
            return json.dumps(PROFILE_JSON)
        if p.startswith("Read this job description"):
            return json.dumps({"job_title": "Dev", "company": "Acme", "must_have": [{"term": "Python", "original": "Python"}]})
        if p.startswith("Write a cover letter"):
            return json.dumps({"letter": "Dear team,\n\nI build APIs.\n\nAda"})
        return json.dumps({"summary": "Python developer.", "experiences": [], "skills": ["Python"]})


def _app(tmp_path):
    fetch = lambda url: FetchResult(ok=True, text="Need Python", reason="", title="Dev", company="Acme")  # noqa: E731
    return create_app(data_dir=tmp_path, model_factory=FakeModel, fetcher=fetch, sync_jobs=True)


def _ready(app, client, url="https://jobs.example.com/1", **form):
    buf = BytesIO()
    d = Document()
    d.add_paragraph("Ada, Python")
    d.add_paragraph("Engineer at Northwind since 2021. Built Python APIs used by three teams and ran on-call for the payments platform.")
    d.save(buf)
    res = client.post("/profile/upload", data={"cv": (BytesIO(buf.getvalue()), "cv.docx")}, content_type="multipart/form-data")
    client.post(res.headers["Location"], data={"action": "send"})  # accept the check-before-sending page
    client.post("/jobs", data={"urls": url})
    run = next(r for r in app.store.list_runs() if r.job_url == url)
    client.post(f"/runs/{run.id}/questions", data=form)
    return app.store.get_run(run.id)


def test_deleting_an_old_cv_does_not_pair_an_old_letter_with_the_new_cv(tmp_path):
    app = _app(tmp_path)
    client = app.test_client()
    run = _ready(app, client, generate_letter="1")        # v1: CV + letter
    client.post(f"/runs/{run.id}/questions", data={})     # v2: CV only
    v1_cv = next(d for d in app.store.list_documents(run.id) if d.kind == "cv" and d.version == 1)
    client.post(f"/documents/{v1_cv.id}/delete")
    after = app.store.get_run(run.id)
    assert after.output_cv_path.endswith("_v2.docx") and after.letter_path == ""


def test_backslash_next_cannot_leave_the_site(tmp_path):
    res = _app(tmp_path).test_client().post("/jobs/bulk", data={"run_id": [], "next": "/\\evil.example.com"})
    assert res.headers["Location"].endswith("/jobs")


def test_saved_api_key_can_be_removed(tmp_path, monkeypatch):
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    client = _app(tmp_path).test_client()
    client.post("/settings", data={"LLM_PROVIDER": "openrouter", "LLM_MODEL": "m", "LLM_API_KEY": "k-123"})
    assert b"Remove the saved key" in client.get("/settings").data
    client.post("/settings", data={"LLM_PROVIDER": "openrouter", "LLM_MODEL": "m", "LLM_API_KEY": "", "clear_api_key": "1"})
    assert effective_settings(tmp_path)[0]["LLM_API_KEY"] == ""


def test_export_names_are_unique_for_same_titles_and_file_names(tmp_path):
    app = _app(tmp_path)
    client = app.test_client()
    _ready(app, client, url="https://jobs.example.com/1")
    _ready(app, client, url="https://jobs.example.com/2")  # same title/company, second upload of cv.docx
    names = zipfile.ZipFile(io.BytesIO(client.get("/settings/export").data)).namelist()
    assert len(names) == len(set(names))
    assert sum(n.startswith("documents/") for n in names) == 2 and sum(n.startswith("source-cvs/") for n in names) == 2


def test_shared_script_is_loaded_on_every_page(tmp_path):
    page = _app(tmp_path).test_client().get("/profile").data
    assert b'src="/static/app.js"' in page and b"form[data-confirm]" not in page
    assert Path(__file__).parents[1].joinpath("cv_maker", "static", "app.js").exists()
