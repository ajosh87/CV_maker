from pathlib import Path

import pytest
from docx import Document

from cv_maker.profile.extract_text import UnreadableCvError, extract_text


def test_extract_docx_paragraphs(tmp_path: Path):
    path = tmp_path / "cv.docx"
    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("Python developer")
    doc.save(path)
    text = extract_text(path)
    assert "Jane Doe" in text
    assert "Python developer" in text


def test_unreadable_extension_raises(tmp_path: Path):
    path = tmp_path / "cv.txt"
    path.write_text("nope")
    with pytest.raises(UnreadableCvError):
        extract_text(path)


def test_pdf_uses_pdfminer(monkeypatch, tmp_path: Path):
    pdf_path = tmp_path / "cv.pdf"
    pdf_path.write_bytes(b"%PDF-1.4 fake")
    monkeypatch.setattr("cv_maker.profile.extract_text.pdf_extract", lambda p: "PDF TEXT")
    text = extract_text(pdf_path)
    assert text == "PDF TEXT"


def test_empty_docx_raises(tmp_path: Path):
    path = tmp_path / "cv.docx"
    doc = Document()
    doc.save(path)
    with pytest.raises(UnreadableCvError, match="empty CV text"):
        extract_text(path)


def test_docx_tables_and_header_are_extracted(tmp_path: Path):
    path = tmp_path / "cv.docx"
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Jane Doe | jane@example.com"
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Senior Engineer"
    table.cell(0, 1).text = "+44 7700 900123"
    doc.add_paragraph("Experience at Acme")
    doc.save(path)
    lines = extract_text(path).splitlines()
    assert lines[0] == "Jane Doe | jane@example.com"
    assert "Senior Engineer | +44 7700 900123" in lines
    assert lines.index("Senior Engineer | +44 7700 900123") < lines.index("Experience at Acme")


def test_docx_content_controls_are_extracted(tmp_path: Path):
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    path = tmp_path / "cv.docx"
    doc = Document()
    doc.add_paragraph("Before")
    sdt = parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtContent><w:p><w:r><w:t>Inside a content control</w:t></w:r></w:p></w:sdtContent></w:sdt>'
    )
    doc.element.body.insert(1, sdt)
    doc.save(path)
    assert "Inside a content control" in extract_text(path)
