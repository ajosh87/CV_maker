"""Personal details stay on this computer: what the LLM receives, and what the user sees and controls."""
import html
import json
import re
from io import BytesIO
from types import SimpleNamespace

from docx import Document

from cv_maker.app import create_app
from cv_maker.jobs.fetch import FetchResult
from cv_maker.pipeline.graph import _prompt_profile, _sign_off
from cv_maker.privacy import (PrivateModel, describe_hidden, fill_missing_contacts, find_in_cv, identity_secrets,
                              leave_out_unused, mask, unmask)


def _profile(**kw):
    return SimpleNamespace(**{"name": "", "email": "", "phone": "", "location": "", "links": {}, **kw})


# ---- masking ----------------------------------------------------------------------------------

def test_your_own_details_are_hidden_in_any_script_and_put_back_exactly():
    for name, place in [("அருண் குமார்", "Chennai, Tamil Nadu"), ("王伟", "北京市朝阳区"), ("محمد الأحمد", "الرياض"),
                        ("Jürgen Müller", "Musterstraße 12, 10115 Berlin"), ("Siobhán O'Neill", "Dublin")]:
        secrets = identity_secrets(_profile(name=name, email="me@example.org", phone="+91 98765 43210", location=place))
        text = f"{name}\n{place}\nme@example.org · +91-98765-43210\nWorked at Acme."
        masked, mapping = mask(text, secrets)
        assert name not in masked and place not in masked and "98765" not in masked and "me@example.org" not in masked
        assert "Worked at Acme." in masked
        assert unmask(masked, mapping) == text


def test_parts_of_your_name_and_capitals_are_hidden_too():
    masked, _ = mask("PRIYA SHARMA\nKind regards,\nPriya\n(Sharma)", identity_secrets(_profile(name="Priya Sharma")))
    assert not re.search(r"(?i)priya|sharma", masked)


def test_accents_match_however_they_were_typed():
    decomposed = "Jürgen Müller"  # what some PDFs produce for "Jürgen Müller"
    masked, _ = mask(f"{decomposed}, Berlin", identity_secrets(_profile(name="Jürgen Müller")))
    assert masked == "[[NAME_1]], Berlin"


def test_contact_formats_are_found_anywhere_but_dates_numbers_and_tech_names_are_not():
    hidden = ["+91 98765 43210", "+44 20 7946 0958", "(555) 123-4567", "090-1234-5678", "06 12 34 56 78", "9876543210",
              "ada@example.com", "https://github.com/ada", "www.ada.dev", "linkedin.com/in/ada-l"]
    kept = ["2019 - 2021", "2016 2017 2018", "01.2019 – 03.2021", "ASP.NET/C#", "Socket.io/Redis", "Node.js",
            "1,200,000 users", "v2.3.1", "35%", "10.0.0.1"]
    for value in hidden:
        assert value not in mask(f"Contact {value} today", [])[0], value
    for value in kept:
        assert value in mask(f"Did {value} today", [])[0], value


def test_restoring_is_json_safe_and_tolerates_how_models_write_placeholders():
    mapping = {"NAME_1": 'Ann "AJ" O\'Neil', "EMAIL_1": "ann@x.io"}
    reply = '{"name": "[[NAME_1]]", "email": "[EMAIL_1]", "note": "{{ name_1 }} via [[ EMAIL_1 ]]", "x": "[[PHONE_9]]"}'
    assert json.loads(unmask(reply, mapping, json_escape=True)) == {
        "name": 'Ann "AJ" O\'Neil', "email": "ann@x.io", "note": 'Ann "AJ" O\'Neil via ann@x.io', "x": ""}


def test_the_model_only_ever_sees_placeholders_and_the_reply_is_restored():
    class Recorder:
        sent = []

        def complete(self, messages, *, json_mode=False):
            self.sent.append(messages[-1]["content"])
            return '{"name": "[[NAME_1]]", "email": "[[EMAIL_1]]", "summary": "Engineer in [[ADDRESS_1]]."}'

    log, inner = [], Recorder()
    secrets = identity_secrets(_profile(name="Ada Lovelace", email="ada@example.com", location="London"))
    model = PrivateModel(inner, lambda: secrets, lambda *entry: log.append(entry))
    reply = model.complete([{"role": "user", "content": "You are tailoring a CV for Ada Lovelace (ada@example.com) in London."}])
    assert json.loads(reply) == {"name": "Ada Lovelace", "email": "ada@example.com", "summary": "Engineer in London."}
    sent = inner.sent[0]
    assert "Ada" not in sent and "ada@example.com" not in sent and "London" not in sent and "[[NAME_1]]" in sent
    assert [entry[:3] for entry in log] == [("Writing a tailored CV", sent, ["ADDRESS_1", "EMAIL_1", "NAME_1"])]
    assert log[0][3]["estimated"] is True and log[0][3]["tokens_in"] > 0  # the stand-in reports no usage


def test_describe_hidden_and_fill_missing_contacts():
    assert describe_hidden(["NAME_1", "EMAIL_1", "EMAIL_2", "DETAIL_1"]) == "name, 2 emails, other detail"
    profile = SimpleNamespace(name="", email="ada@example.com", phone="", location="")
    fill_missing_contacts(profile, {"NAME_1": "Ada\nLovelace", "PHONE_1": "+44 20 7946 0958", "EMAIL_1": "other@example.com"})
    assert (profile.name, profile.email, profile.phone) == ("Ada Lovelace", "ada@example.com", "+44 20 7946 0958")


# ---- reading a CV -----------------------------------------------------------------------------

def test_a_cv_header_is_understood_in_several_languages_and_layouts():
    cases = {
        "Priya Sharma\nSenior Data Engineer\npriya@gmail.com | +91 98765 43210 | No. 12, 3rd Cross, Indiranagar, Bengaluru 560038\n"
        "EXPERIENCE\nData engineer at Infosys": {("NAME", "Priya Sharma"), ("ADDRESS", "No. 12, 3rd Cross, Indiranagar, Bengaluru 560038")},
        "Jürgen Müller\nMusterstraße 12, 10115 Berlin\nTel.: +49 30 1234567 · juergen@example.de\nBerufserfahrung":
            {("NAME", "Jürgen Müller"), ("ADDRESS", "Musterstraße 12, 10115 Berlin")},
        "王伟\n电话: 138 0013 8000\n北京市朝阳区建国路88号\n工作经历": {("NAME", "王伟"), ("ADDRESS", "北京市朝阳区建国路88号")},
        "محمد الأحمد\nالرياض، المملكة العربية السعودية\n+966 50 123 4567\nExperience":
            {("NAME", "محمد الأحمد"), ("ADDRESS", "الرياض، المملكة العربية السعودية")},
        "அருண் குமார்\nChennai, Tamil Nadu\narun@example.in\nSkills": {("NAME", "அருண் குமார்"), ("ADDRESS", "Chennai, Tamil Nadu")},
        "CURRICULUM VITAE\nSenior Software Engineer\nArjun Nair\narjun@example.com | Kochi, India\nSummary":
            {("NAME", "Arjun Nair"), ("ADDRESS", "Kochi, India")},
        "Name: Lena Fischer\nAddress:\nHauptstraße 5\n80331 München\nSkills":
            {("NAME", "Lena Fischer"), ("ADDRESS", "Hauptstraße 5"), ("ADDRESS", "80331 München")},
        "New Delhi, India | +91 98765 43210\nRohan Mehta, MBA, PMP\nAWS Certified 2022 | Data Platforms\nExperience":
            {("ADDRESS", "New Delhi, India"), ("NAME", "Rohan Mehta")},
    }
    for text, expected in cases.items():
        assert {(s.kind, s.text) for s in find_in_cv(text)} == expected, text


def test_parts_the_app_never_uses_are_left_out():
    text = ("Ada Lovelace\nEXPERIENCE\nEngineer at Northwind\nPERSONAL DETAILS\nDate of Birth: 10/12/1990\n"
            "Marital Status: Single\nPan-India rollout: led 5 teams\nLanguages: English, Tamil\nREFERENCES\n"
            "Dr. Anil Kumar, IISc\nanil@iisc.ac.in\n\nDECLARATION\nI declare the above is true.\n")
    kept, left_out = leave_out_unused(text)
    assert "Date of Birth" not in kept and "Single" not in kept and "Anil" not in kept and "declare" not in kept
    assert "Pan-India rollout: led 5 teams" in kept and "Languages: English, Tamil" in kept and "Engineer at Northwind" in kept
    assert left_out == ["Date of birth", "Marital status", "References (2 lines)", "Declaration (1 line)"]


# ---- writing ----------------------------------------------------------------------------------

def test_writing_prompts_never_include_contact_details():
    profile = {"name": "Ada", "email": "a@x.io", "phone": "1", "location": "London", "links": {"github": "g"}, "target_role": "x",
               "experiences": [{"company": "Northwind", "title": "Engineer", "location": "", "start": "2021", "end": None,
                                "current": True, "bullets": ["b"], "confirmed": True}],
               "education": [], "skills": [{"name": "Python", "source": "cv"}], "extras": {"projects": ["P"], "dob": "1990"}}
    data = json.loads(_prompt_profile(profile))
    assert set(data) == {"experiences", "education", "skills", "extras"} and data["extras"] == {"projects": ["P"]}


def test_the_app_signs_the_letter_itself():
    assert _sign_off("Dear team,\n\nI build APIs.\n\nKind regards,", "Ada King") == "Dear team,\n\nI build APIs.\n\nKind regards,\nAda King"
    assert _sign_off("Dear team,\n\nI build APIs.\n\nBest regards,\n[Your Name]", "Ada King").endswith("\n\nBest regards,\nAda King")
    assert _sign_off("Dear team,\n\nI build APIs.", "Ada King").endswith("I build APIs.\n\nKind regards,\nAda King")
    assert _sign_off("Dear team,\n\nI build APIs.\n\nSincerely\nAda King", "Ada King").endswith("\n\nSincerely\nAda King")
    assert _sign_off("Dear team,\n\nI build APIs.", "") == "Dear team,\n\nI build APIs."


# ---- in the app -------------------------------------------------------------------------------

PROFILE = {"name": "Ada Lovelace", "email": "ada@example.com", "phone": "+44 20 7946 0958", "location": "London",
           "links": {"github": "https://github.com/ada"},
           "experiences": [{"company": "Northwind", "title": "Engineer", "location": "London", "start": "2021", "end": None,
                            "current": True, "bullets": ["Built Python APIs"]}],
           "education": [], "skills": ["Python"]}


class RecordingModel:
    """Answers like a model would, and keeps every prompt it was sent."""
    prompts: list = []

    def complete(self, messages, *, json_mode=False):
        prompt = messages[-1]["content"]
        RecordingModel.prompts.append(prompt)
        if prompt.startswith("Extract a structured profile"):
            reply = json.dumps(PROFILE)
            for key, value in (("NAME_1", "Ada Lovelace"), ("EMAIL_1", "ada@example.com"), ("DETAIL_1", "Northwind")):
                if f"[[{key}]]" in prompt:
                    reply = reply.replace(value, f"[[{key}]]")  # a real model copies the placeholder it was given
            return reply
        if prompt.startswith("Read this job description"):
            return json.dumps({"job_title": "Platform Engineer", "company": "Acme", "must_have": [{"term": "Python", "original": "Python"}]})
        if prompt.startswith("You are tailoring"):
            return json.dumps({"summary": "Python engineer.", "skills": ["Python"],
                               "experiences": [{"company": "Northwind", "title": "Engineer", "bullets": ["Built Python APIs"]}]})
        return json.dumps({"letter": "Dear team,\n\nI build Python APIs.\n\nKind regards,"})


def _app(tmp_path):
    RecordingModel.prompts = []
    fetch = lambda url: FetchResult(ok=True, text="Acme is hiring. Need Python. Questions to jane@acme.example or +1 415 555 0100.",  # noqa: E731
                                    reason="", title="Platform Engineer", company="Acme")
    return create_app(data_dir=tmp_path, model_factory=RecordingModel, fetcher=fetch, sync_jobs=True)


def _cv() -> bytes:
    buf = BytesIO()
    doc = Document()
    for line in ["Ada Lovelace", "ada@example.com | +44 20 7946 0958 | 12 Baker Street, London NW1 6XE",
                 "github.com/ada", "EXPERIENCE", "Engineer, Northwind (2021 – Present)", "Built Python APIs",
                 "Date of Birth: 10 Dec 1990", "REFERENCES", "Charles Babbage, Analytical Engines Ltd, +44 20 7000 0000"]:
        doc.add_paragraph(line)
    doc.save(buf)
    return buf.getvalue()


def _upload(client):
    return client.post("/profile/upload", data={"cv": (BytesIO(_cv()), "cv.docx")}, content_type="multipart/form-data")


def _preview(page: str) -> str:
    return html.unescape(re.search(r'<div class="jd-text sent-text">(.*?)</div>', page, re.S).group(1))


def _ticked(page: str) -> list[tuple[str, str]]:
    """(index, text) of the details ticked to be hidden on the check page."""
    return re.findall(r'name="hide" value="(\d+)" checked>\s*<span class="kind">\w+</span>\s*<span class="break">([^<]+)</span>', page)


PERSONAL = ["Ada", "Lovelace", "ada@example.com", "7946", "Baker Street", "github.com/ada", "1990", "Babbage"]


def test_a_cv_is_only_sent_after_the_check_and_without_personal_details(tmp_path):
    app = _app(tmp_path)
    client = app.test_client()
    res = _upload(client)
    assert res.headers["Location"].endswith("/review") and RecordingModel.prompts == []  # nothing sent yet

    page = client.get(res.headers["Location"]).data.decode()
    assert "Check before sending" in page and "Hidden from the LLM" in page
    assert "12 Baker Street, London NW1 6XE" in page and "Date of birth" in page and "References (1 line)" in page
    preview = _preview(page)
    assert '<mark class="ph">[[NAME_1]]</mark>' in preview and "Built Python APIs" in preview
    assert not [p for p in PERSONAL if p in preview]

    client.post(res.headers["Location"], data={"action": "send"})
    assert len(RecordingModel.prompts) == 1 and RecordingModel.prompts[0].startswith("Extract a structured profile")
    assert not [p for p in PERSONAL if p in RecordingModel.prompts[0]]
    assert re.sub(r"</?mark[^>]*>", "", _preview(page)) in RecordingModel.prompts[0]  # the preview is what was sent
    profile = app.store.get_profile()
    assert (profile.name, profile.email) == ("Ada Lovelace", "ada@example.com")  # put back on this computer


def test_you_choose_what_is_hidden(tmp_path):
    app = _app(tmp_path)
    client = app.test_client()
    review = _upload(client).headers["Location"]
    keep = [i for i, text in _ticked(client.get(review).data.decode()) if text != "ada@example.com"]
    client.post(review, data={"items_shown": "1", "hide": keep, "also": "Northwind", "unused_shown": "1", "leave_out": "1",
                              "action": "preview"})
    page = client.get(review).data.decode()
    preview = _preview(page)
    assert "ada@example.com" in preview and "Northwind" not in preview and "[[DETAIL_1]]" in preview
    assert "Date of" not in preview  # still left out

    # Send with the list as now shown, and this time keep the parts the app doesn't use.
    client.post(review, data={"items_shown": "1", "hide": [i for i, _ in _ticked(page)], "unused_shown": "1", "action": "send"})
    sent = RecordingModel.prompts[0]
    assert "ada@example.com" in sent and "Northwind" not in sent and "Date of Birth" in sent
    assert app.store.get_profile().experiences[0].company == "Northwind"  # put back in the answer


def test_cancel_deletes_the_cv_without_sending_it(tmp_path):
    app = _app(tmp_path)
    client = app.test_client()
    review = _upload(client).headers["Location"]
    client.post(review.replace("/review", "/delete"))
    assert app.store.list_uploads() == [] and RecordingModel.prompts == []
    assert client.get(review).status_code == 404


def test_every_request_is_masked_recorded_and_signed_locally(tmp_path):
    app = _app(tmp_path)
    client = app.test_client()
    client.post(_upload(client).headers["Location"], data={"action": "send"})
    client.post("/settings/privacy", data={"never_send": "Acme\n  acme  \n"})
    assert app.store.get_never_send() == ["Acme"]
    client.post("/jobs", data={"urls": "https://jobs.example.com/1"})
    run = app.store.list_runs()[0]
    client.post(f"/runs/{run.id}/questions", data={"generate_letter": "1"})
    run = app.store.get_run(run.id)
    assert run.status == "ready" and run.company == "Acme"

    reading_job, planning_cv, writing_cv, writing_letter, prep_notes = RecordingModel.prompts[1:]
    assert "jane@acme.example" not in reading_job and "555 0100" not in reading_job and "Acme" not in reading_job
    assert planning_cv.startswith("You are planning how to tailor a CV")
    for prompt in (planning_cv, writing_cv, writing_letter, prep_notes):
        assert not [p for p in ["Ada", "Lovelace", "ada@example.com", "7946", "github.com/ada", "London"] if p in prompt]
        assert "Built Python APIs" in prompt  # the career facts the LLM needs are still there
    assert prep_notes.startswith("You are helping a candidate prepare for interviews")
    letter = [p.text for p in Document(run.letter_path).paragraphs]
    assert letter[0] == "Ada Lovelace" and letter[-1] == "Kind regards,\nAda Lovelace"

    page = client.get("/settings/sent").data.decode()
    assert page.count('class="version sent-call"') == 6 and "Writing interview prep notes" in page
    assert "Planning the CV for this job" in page
    assert "Reading your CV" in page and "Writing a cover letter" in page and '<mark class="ph">' in page
    assert "Hidden: name, email, phone number, link, address." in page
    client.post("/settings/sent/clear")
    assert app.store.list_llm_calls() == []


def test_delete_all_removes_the_privacy_list_and_the_record(tmp_path):
    app = _app(tmp_path)
    client = app.test_client()
    client.post(_upload(client).headers["Location"], data={"action": "send"})
    client.post("/settings/privacy", data={"never_send": "Northwind"})
    client.post("/settings/delete-all", data={"confirm": "DELETE"})
    assert app.store.get_never_send() == [] and app.store.list_llm_calls() == []


def test_settings_explain_privacy(tmp_path):
    page = _app(tmp_path).test_client().get("/settings").data.decode()
    assert 'id="privacy"' in page and "Also never send" in page and "See exactly what was sent" in page
