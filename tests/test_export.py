from pathlib import Path
from docx import Document
from cv_maker.honesty import CvDocument
from cv_maker.export.docx_cv import ExportError, write_cv_docx
from cv_maker.export.letter import write_letter_docx


def test_cv_headings_no_tables(tmp_path: Path):
    path = tmp_path / "cv.docx"
    write_cv_docx(
        CvDocument(
            name="Ada",
            summary="Engineer",
            experiences=[],
            education=[],
            skills=["Python"],
        ),
        path,
    )
    d = Document(path)
    heading_texts = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert "Experience" in heading_texts
    assert "Education" in heading_texts
    assert "Skills" in heading_texts
    assert d.tables == []
    assert d.sections[0].header.paragraphs[0].text.strip() == ""
    fonts = {run.font.name for p in d.paragraphs for run in p.runs if run.font.name}
    assert fonts <= {"Calibri", None} or "Calibri" in fonts


def test_letter_written(tmp_path: Path):
    path = tmp_path / "letter.docx"
    write_letter_docx("Dear hiring manager, I am Ada.", path)
    assert path.exists()
    assert "Ada" in "\n".join(p.text for p in Document(path).paragraphs)


def test_export_error_no_fake_file(tmp_path: Path):
    path = tmp_path / "cv.docx"
    try:
        write_cv_docx(
            CvDocument(
                name="Ada",
                summary="Engineer",
                experiences=[object()],
                education=[],
                skills=[],
            ),
            path,
        )
    except ExportError:
        pass
    assert not path.exists()


def test_cv_includes_contact_dates_and_education(tmp_path: Path):
    from cv_maker.models import Education, Experience

    path = tmp_path / "cv.docx"
    write_cv_docx(
        CvDocument(
            name="Priya Raman",
            summary="Backend engineer.",
            experiences=[Experience("Northwind", "Senior Engineer", "London", "Mar 2021", None, True, ["Built a ledger"], True)],
            education=[Education("University of Leeds", "BEng", "Software Engineering", "2010 - 2014")],
            skills=["Python", "Go"],
            contact=["priya@example.com", "+44 7700 900123"],
            headline="Senior Engineer",
            extras={"certifications": ["AWS SA Associate"]},
        ),
        path,
    )
    text = "\n".join(p.text for p in Document(path).paragraphs)
    for expected in (
        "Priya Raman",
        "priya@example.com | +44 7700 900123",
        "Senior Engineer — Northwind",
        "London | Mar 2021 – Present",
        "BEng, Software Engineering — University of Leeds (2010 - 2014)",
        "AWS SA Associate",
    ):
        assert expected in text
    assert Document(path).tables == []


def test_letter_has_paragraphs_and_signature_block(tmp_path: Path):
    path = tmp_path / "letter.docx"
    write_letter_docx("Dear team,\n\nI build ledgers.\n\nPriya", path, name="Priya", contact=["p@x.com"])
    paragraphs = [p.text for p in Document(path).paragraphs]
    assert paragraphs[:2] == ["Priya", "p@x.com"]
    assert paragraphs[-3:] == ["Dear team,", "I build ledgers.", "Priya"]
