import json

from docx import Document

from cv_maker.llm.chat import LLMError
from cv_maker.models import Experience, Profile, Skill
from cv_maker.pipeline.graph import Pipeline
from cv_maker.store import Store

REQUIREMENTS = {
    "job_title": "Platform Engineer", "company": "Acme",
    "must_have": [{"term": "Python", "original": "Python"}, {"term": "Kubernetes", "original": "Kubernetes in production"}],
    "nice_to_have": [], "tools": [], "seniority": "", "domain": "", "education": "", "language": "", "work_auth": "",
}
REWRITE = {
    "summary": "Python engineer who ships reliable services. Kubernetes expert.",
    "experiences": [{"company": "Northwind", "title": "Engineer", "bullets": ["Built Python APIs serving 3 teams"]}],
    "skills": ["Python", "Kubernetes"],
}


class FakeModel:
    def __init__(self, fail_on=None, fenced=False):
        self.calls = []
        self.fail_on = fail_on
        self.fenced = fenced

    def complete(self, messages, *, json_mode=False):
        prompt = messages[-1]["content"]
        kind = ("requirements" if prompt.startswith("Read this job description") else
                "rewrite" if prompt.startswith("You are tailoring") else
                "letter" if prompt.startswith("Write a cover letter") else "other")
        self.calls.append(kind)
        if kind == self.fail_on:
            raise LLMError("OpenRouter rejected the API key (HTTP 401). Update it on the Settings page.", status=401)
        body = {"requirements": REQUIREMENTS, "rewrite": REWRITE,
                "letter": {"letter": "Dear team,\n\nI build Python APIs. I run Kubernetes.\n\nAda"}}.get(kind, {})
        text = json.dumps(body)
        return f"Here you go:\n```json\n{text}\n```" if self.fenced else text


def _setup(tmp_path, model):
    store = Store(tmp_path / "db.sqlite")
    store.save_profile(Profile(
        name="Ada", email="ada@x.com", phone="", location="London", links={}, target_role=None,
        experiences=[Experience("Northwind", "Engineer", "London", "2021", None, True, ["Built APIs serving 3 teams"], True)],
        education=[], skills=[Skill("Python", "cv")], extras={},
    ))
    run = store.create_run(job_url="https://x")
    run.jd_text = "Need Python and Kubernetes in production"
    store.update_run(run)
    return store, Pipeline(store, lambda: model, tmp_path), run


def _answer(store, run_id, choice):
    run = store.get_run(run_id)
    run.answers = [{"term": "Kubernetes", "skipped": choice == "skip", "text": "yes" if choice == "yes" else "", "add_role": None}]
    store.update_run(run)


def test_skip_kubernetes_absent_from_export(tmp_path):
    store, pipeline, run = _setup(tmp_path, FakeModel())
    analyzed = pipeline.analyze(run.id)
    assert analyzed.status == "needs_answers"
    assert [q["term"] for q in analyzed.questions if not q.get("optional")] == ["Kubernetes"]
    assert (analyzed.title, analyzed.company) == ("Platform Engineer", "Acme")

    _answer(store, run.id, "skip")
    done = pipeline.generate(run.id)
    assert done.status == "ready", done.error
    assert done.draft["skills"] == ["Python"]
    assert "Kubernetes" not in done.draft["summary"]
    assert done.draft["summary"] == "Python engineer who ships reliable services."
    text = "\n".join(p.text for p in Document(done.output_cv_path).paragraphs)
    assert "Kubernetes" not in text and "ada@x.com" in text


def test_yes_answer_reaches_profile_and_cv(tmp_path):
    store, pipeline, run = _setup(tmp_path, FakeModel())
    pipeline.analyze(run.id)
    _answer(store, run.id, "yes")
    done = pipeline.generate(run.id)
    assert "Kubernetes" in done.draft["skills"]
    assert any(s.name == "Kubernetes" and s.source == "clarification" for s in store.get_profile().skills)


def test_llm_failure_stops_the_graph_and_is_recorded(tmp_path):
    model = FakeModel(fail_on="requirements")
    store, pipeline, run = _setup(tmp_path, model)
    failed = pipeline.analyze(run.id)
    assert failed.status == "failed"
    assert failed.failed_stage == "analyze"
    assert "rejected the API key" in failed.error
    assert model.calls == ["requirements"]  # nothing ran after the failing node
    assert failed.step == ""


def test_retry_after_failure_works(tmp_path):
    model = FakeModel(fail_on="requirements")
    store, pipeline, run = _setup(tmp_path, model)
    pipeline.analyze(run.id)
    model.fail_on = None
    retried = pipeline.analyze(run.id)
    assert retried.status == "needs_answers" and retried.error == ""


def test_missing_profile_is_explained(tmp_path):
    store = Store(tmp_path / "db.sqlite")
    run = store.create_run(job_url="https://x")
    run.jd_text = "Need Python"
    store.update_run(run)
    failed = Pipeline(store, lambda: FakeModel(), tmp_path).analyze(run.id)
    assert failed.status == "failed" and "Upload your CV" in failed.error


def test_letter_failure_is_only_a_warning(tmp_path):
    store, pipeline, run = _setup(tmp_path, FakeModel(fail_on="letter"))
    pipeline.analyze(run.id)
    _answer(store, run.id, "skip")
    run = store.get_run(run.id)
    run.generate_letter = True
    store.update_run(run)
    done = pipeline.generate(run.id)
    assert done.status == "ready"
    assert done.letter_path == ""
    assert "Cover letter not created" in done.warning


def test_letter_is_generated_and_filtered(tmp_path):
    store, pipeline, run = _setup(tmp_path, FakeModel())
    pipeline.analyze(run.id)
    _answer(store, run.id, "skip")
    run = store.get_run(run.id)
    run.generate_letter = True
    store.update_run(run)
    done = pipeline.generate(run.id)
    text = "\n".join(p.text for p in Document(done.letter_path).paragraphs)
    assert "I build Python APIs." in text
    assert "Kubernetes" not in text


def test_fenced_json_replies_are_accepted(tmp_path):
    store, pipeline, run = _setup(tmp_path, FakeModel(fenced=True))
    assert pipeline.analyze(run.id).status == "needs_answers"
    _answer(store, run.id, "skip")
    assert pipeline.generate(run.id).status == "ready"


def test_rewrite_failure_keeps_run_retryable_at_generate_stage(tmp_path):
    store, pipeline, run = _setup(tmp_path, FakeModel(fail_on="rewrite"))
    pipeline.analyze(run.id)
    _answer(store, run.id, "skip")
    failed = pipeline.generate(run.id)
    assert (failed.status, failed.failed_stage) == ("failed", "generate")
    assert failed.questions  # analysis results survive for the retry
