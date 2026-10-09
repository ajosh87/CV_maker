import io
import json
import re
import zipfile
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from cv_maker.app import create_app
from cv_maker.jobs.fetch import FetchResult
from cv_maker.llm.chat import ProviderUnsetError

PROFILE = {
    "name": "Ada Lovelace", "email": "ada@example.com", "phone": "+44 1", "location": "London", "links": {},
    "experiences": [{"company": "Northwind", "title": "Engineer", "location": "London", "start": "2021", "end": None,
                     "current": True, "bullets": ["Built Python APIs"]}],
    "education": [], "skills": ["Python"],
}
REQUIREMENTS = {
    "job_title": "Platform Engineer", "company": "Acme",
    "must_have": [{"term": "Python", "original": "Python"}, {"term": "Kubernetes", "original": "Kubernetes"}],
    "nice_to_have": [],
}


class FakeModel:
    def complete(self, messages, *, json_mode=False):
        prompt = messages[-1]["content"]
        if prompt.startswith("Extract a structured profile"):
            return json.dumps(PROFILE)
        if prompt.startswith("Read this job description"):
            return json.dumps(REQUIREMENTS)
        if prompt.startswith("You are tailoring"):
            return json.dumps({"summary": "Python engineer. Kubernetes guru.", "skills": ["Python", "Kubernetes"],
                               "experiences": [{"company": "Northwind", "title": "Engineer", "bullets": ["Built Python APIs for Acme-style platforms"]}]})
        if prompt.startswith("Write a cover letter"):
            return json.dumps({"letter": "Dear team,\n\nI build Python APIs.\n\nAda Lovelace"})
        return "{}"


def fake_fetch(url):
    if url.endswith("/fail"):
        return FetchResult(ok=False, text="", reason="The site blocked automated access (HTTP 403).")
    return FetchResult(ok=True, text="Need Python and Kubernetes", reason="", title="Platform Engineer", company="Acme")


def _docx_bytes(extra_chars: int = 0) -> bytes:
    buf = BytesIO()
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Ada Lovelace"
    table.cell(0, 1).text = "ada@example.com"
    doc.add_paragraph("Engineer at Northwind. Python.")
    doc.add_paragraph("Engineer at Northwind since 2021. Built Python APIs used by three teams and ran on-call for the payments platform.")  # a real CV has more text than a scanned image would
    if extra_chars:
        doc.add_paragraph("x " * (extra_chars // 2))
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def app(tmp_path):
    return create_app(data_dir=tmp_path, model_factory=FakeModel, fetcher=fake_fetch, sync_jobs=True)


def _upload(client, data=None, name="cv.docx", mode="merge", send=True):
    """Upload a CV and, like a user accepting the check-before-sending page, send it."""
    res = client.post("/profile/upload", data={"cv": (BytesIO(data or _docx_bytes()), name), "mode": mode},
                      content_type="multipart/form-data")
    if send and "/review" in res.headers.get("Location", ""):
        client.post(res.headers["Location"], data={"action": "send"})
    return res


def _session_cookie(res) -> str:
    header = next((h for h in res.headers.getlist("Set-Cookie") if h.startswith("session=")), "")
    return header.split(";", 1)[0]


def _ready_job(app, client, url="https://jobs.example.com/ok1", **answers):
    _upload(client)
    client.post("/jobs", data={"urls": url})
    run = next(r for r in app.store.list_runs() if r.job_url == url)
    client.post(f"/runs/{run.id}/questions", data={"answer_0": "skip", **answers})
    return app.store.get_run(run.id)


# ---- navigation & first run ----

def test_home_sends_new_users_to_profile_and_returning_users_to_jobs(app):
    client = app.test_client()
    assert client.get("/").headers["Location"].endswith("/profile")
    _upload(client)
    assert client.get("/").headers["Location"].endswith("/jobs")


def test_first_run_profile_page_explains_how_it_works(app):
    page = app.test_client().get("/profile").data
    assert b"Drop your CV here" in page and b'class="how' in page
    assert b"sent to the LLM you choose" in page  # transparency before the first upload


def test_old_addresses_redirect(app):
    client = app.test_client()
    assert client.get("/library?kind=cv").headers["Location"].endswith("/documents?kind=cv")
    assert client.get("/library/jobs?bucket=action").headers["Location"].endswith("/jobs?bucket=needs")
    assert client.get("/result").headers["Location"].endswith("/jobs")


def test_single_navigation_with_four_sections(app):
    page = app.test_client().get("/profile").data.decode()
    nav = re.search(r'<nav class="main-nav".*?</nav>', page, re.S).group(0)
    assert re.findall(r">(\w+)</a>", nav) == ["Jobs", "Documents", "Profile", "Settings"]
    assert 'aria-current="page"' in nav and "STEP" not in page.upper().replace("STEPS", "")


# ---- profile ----

def test_bad_cv_creates_no_runs_and_is_not_kept(app, tmp_path):
    client = app.test_client()
    res = client.post("/", data={"cv": (BytesIO(b"not-a-cv"), "cv.txt")}, content_type="multipart/form-data")
    assert res.status_code == 400
    assert b"Upload a PDF or DOCX" in res.data
    assert app.store.list_runs() == []
    assert list((tmp_path / "uploads").iterdir()) == []


def test_upload_parses_profile_and_keeps_cookie_small(app):
    client = app.test_client()
    res = _upload(client, _docx_bytes(extra_chars=30000))
    assert res.status_code == 302 and res.headers["Location"].endswith("/review")
    assert len(_session_cookie(res)) < 500  # CV text lives in SQLite, not the cookie
    page = client.get("/profile").data
    assert b"Ada Lovelace" in page and b"Engineer \xe2\x80\x94 Northwind" in page
    assert b"Built Python APIs" in page  # the full saved profile, not a summary
    assert b"Source CVs" in page and b"cv.docx" in page
    assert b"Edit profile" in page
    assert app.store.get_profile().email == "ada@example.com"


def test_profile_failure_is_explained_with_retry(tmp_path):
    def unset():
        raise ProviderUnsetError("No LLM configured: set provider in .env or on the Settings page")

    app = create_app(data_dir=tmp_path, model_factory=unset, fetcher=fake_fetch, sync_jobs=True)
    client = app.test_client()
    _upload(client)
    page = client.get("/profile").data
    assert b"No LLM configured" in page
    assert b"Retry" in page and b"/settings" in page


def test_edit_profile_changes_facts_and_marks_new_skills(app):
    client = app.test_client()
    _upload(client)
    form = client.get("/profile/edit").data
    assert b'name="exp-0-title" value="Engineer"' in form and b'name="exp-1-title" value=""' in form

    res = client.post("/profile/edit", data={
        "name": "Ada King", "email": "ada@example.com", "phone": "", "location": "London", "links": "github: github.com/ada",
        "exp-0-title": "Senior Engineer", "exp-0-company": "Northwind", "exp-0-start": "2021", "exp-0-current": "1",
        "exp-0-bullets": "Built Python APIs\nLed a team of 4",
        "exp-1-title": "Analyst", "exp-1-company": "Babbage & Co", "exp-1-start": "2018", "exp-1-end": "2020", "exp-1-bullets": "",
        "skills": "Python\nGo", "extra-certifications": "AWS SA",
    })
    assert res.status_code == 302
    profile = app.store.get_profile()
    assert profile.name == "Ada King" and profile.links == {"github": "github.com/ada"}
    assert [(e.title, e.current) for e in profile.experiences] == [("Senior Engineer", True), ("Analyst", False)]
    assert profile.experiences[0].bullets == ["Built Python APIs", "Led a team of 4"]
    assert {s.name: s.source for s in profile.skills} == {"Python": "cv", "Go": "user"}
    assert profile.extras["certifications"] == ["AWS SA"]

    client.post("/profile/edit", data={"name": "Ada King", "exp-0-title": "Senior Engineer", "exp-0-company": "Northwind",
                                       "exp-0-remove": "1", "skills": "Python"})
    assert app.store.get_profile().experiences == []


def test_delete_source_cv_keeps_profile(app, tmp_path):
    client = app.test_client()
    _upload(client)
    upload = app.store.list_uploads()[0]
    assert Path(upload.path).exists()
    client.post(f"/uploads/{upload.id}/delete")
    assert app.store.list_uploads() == [] and not Path(upload.path).exists()
    assert app.store.get_profile() is not None


# ---- jobs ----

def test_jobs_page_without_profile_points_to_profile(app):
    page = app.test_client().get("/jobs").data
    assert b"Start with your CV" in page and b"Add your CV" in page
    assert b'id="jobs-form"' not in page


def test_jobs_page_first_run_shows_add_form_and_steps(app):
    client = app.test_client()
    _upload(client)
    page = client.get("/jobs").data.decode()
    board, right = page.split('id="jobs-dock"')
    assert 'class="how' in board and 'id="jobs-form"' in right  # how it works with the list; adding a job beside it


def test_empty_or_invalid_batch_is_rejected(app):
    client = app.test_client()
    _upload(client)
    res = client.post("/jobs", data={"urls": "not a url\n\n"})
    assert res.status_code == 400
    assert b"Add at least one job link" in res.data and b"not a url" in res.data
    assert app.store.list_runs() == []


def test_full_flow_paste_answer_letter_download(app):
    client = app.test_client()
    _upload(client)
    urls = "https://jobs.example.com/ok1\nhttps://jobs.example.com/fail\nhttps://jobs.example.com/ok2"
    res = client.post("/jobs", data={"urls": urls})
    assert res.headers["Location"].endswith("/jobs")

    runs = {r.job_url: r for r in app.store.list_runs()}
    assert runs["https://jobs.example.com/fail"].status == "needs_paste"
    assert runs["https://jobs.example.com/ok1"].status == "needs_answers"
    board = client.get("/jobs").data
    assert b"Paste description" in board and b"Answer 1 question" in board

    failed = runs["https://jobs.example.com/fail"]
    detail = client.get(f"/runs/{failed.id}").data
    assert b"blocked automated access" in detail and b'id="paste"' in detail
    client.post(f"/runs/{failed.id}/paste", data={"jd_text": "Platform Engineer. Need Python and Kubernetes, hybrid."})
    assert app.store.get_run(failed.id).status == "needs_answers"

    first = runs["https://jobs.example.com/ok1"]
    page = client.get(f"/runs/{first.id}/questions").data
    assert b"Platform Engineer" in page and b"job 1 of 3" in page
    assert b'name="answer_0" value="intermediate"' in page and b"Kubernetes" in page  # a level, not just yes / no

    res = client.post(f"/runs/{first.id}/questions", data={"answer_0": "intermediate", "detail_0": "Ran clusters at Northwind",
                                                           "generate_letter": "1"})
    assert "/questions" in res.headers["Location"]  # sent on to the next job needing answers
    for run in app.store.list_runs():
        if run.status == "needs_answers":
            client.post(f"/runs/{run.id}/questions", data={"answer_0": "skip"})

    done = {r.job_url: r for r in app.store.list_runs()}
    assert {r.status for r in done.values()} == {"ready"}
    assert "Kubernetes" in done["https://jobs.example.com/ok1"].draft["skills"]  # confirmed by the user
    assert done["https://jobs.example.com/ok1"].letter_path

    with client.get(f"/download/{first.id}/cv") as cv:
        assert cv.status_code == 200
        assert re.search(r'filename="?Ada_Lovelace_CV_v1_-_Acme\.docx', cv.headers["Content-Disposition"])
    with client.get(f"/download/{first.id}/letter") as letter:
        assert letter.status_code == 200
    board, right = client.get("/jobs").data.decode().split('id="jobs-dock"')
    assert board.count("Download CV") == 3 and "Download CV" not in right  # said once: in the list, not beside it too


def test_skipped_requirement_never_reaches_the_cv(app):
    client = app.test_client()
    done = _ready_job(app, client, answer_0="no")
    assert done.status == "ready"
    assert done.draft["skills"] == ["Python"]
    assert "Kubernetes" not in done.draft["summary"]


def test_paste_only_job(app):
    client = app.test_client()
    _upload(client)
    client.post("/jobs", data={"urls": "", "pasted_jd": "Too short", "pasted_title": "Infra role"})
    assert app.store.list_runs() == []  # not a job description: nothing added, and the page says why
    res = client.post("/jobs", data={"urls": "", "pasted_jd": "Infra role. Need Python and Kubernetes, hybrid in London.",
                                     "pasted_title": "Infra role"}, follow_redirects=True)
    run = app.store.list_runs()[0]
    assert run.job_url is None and run.status == "needs_answers" and run.title == "Infra role"
    assert "That description is short" in res.data.decode()


def test_jobs_buckets_and_bulk_archive_and_delete(app):
    client = app.test_client()
    ready = _ready_job(app, client)
    client.post("/jobs", data={"urls": "https://jobs.example.com/fail"})
    stuck = next(r for r in app.store.list_runs() if r.job_url.endswith("/fail"))

    page = client.get("/jobs").data.decode()
    assert re.search(r'Ready<span class="count">1</span>', page) and re.search(r'Needs you<span class="count">1</span>', page)
    assert client.get("/jobs?bucket=ready").data.count(b'name="run_id"') == 1

    res = client.post("/jobs/bulk", data={"run_id": [stuck.id], "action": "archive", "next": "/jobs?bucket=needs"})
    assert res.headers["Location"].endswith("/jobs?bucket=needs") and app.store.get_run(stuck.id).archived
    assert b"under <a" in client.get("/jobs?bucket=needs").data  # points to Archived instead of a dead end

    client.post("/jobs/bulk", data={"run_id": [stuck.id], "action": "unarchive"})
    assert not app.store.get_run(stuck.id).archived

    cv_path = Path(ready.output_cv_path)
    client.post("/jobs/bulk", data={"run_id": [ready.id, stuck.id], "action": "delete"})
    assert app.store.list_runs() == [] and app.store.list_documents() == [] and not cv_path.exists()


def test_retry_returns_json_for_xhr(app):
    client = app.test_client()
    _upload(client)
    run = app.store.create_run(job_url="https://jobs.example.com/x")
    body = client.post(f"/retry/{run.id}", headers={"X-Requested-With": "XMLHttpRequest"}).get_json()
    assert body["status"] == "needs_answers"
    assert body["jd_text"] == "Need Python and Kubernetes"


def test_retry_after_llm_failure_keeps_pasted_description(app):
    client = app.test_client()
    _upload(client)
    run = app.store.create_run(job_url="https://jobs.example.com/fail")
    run.jd_text, run.status, run.failed_stage, run.error = "Pasted by hand: Need Python", "failed", "analyze", "boom"
    app.store.update_run(run)
    client.post(f"/runs/{run.id}/retry")
    after = app.store.get_run(run.id)
    assert after.jd_text == "Pasted by hand: Need Python"
    assert after.status == "needs_answers"


def test_downloads_for_missing_files_do_not_crash(app):
    client = app.test_client()
    run = app.store.create_run(job_url="https://jobs.example.com/x")
    assert client.get(f"/download/{run.id}/cv").status_code == 302
    assert client.get("/download/does-not-exist/cv").status_code == 404
    assert client.post("/retry/does-not-exist").status_code == 404


def test_api_runs_reports_status(app):
    client = app.test_client()
    run = app.store.create_run(job_url="https://jobs.example.com/x", status="fetching")
    body = client.get(f"/api/runs?ids={run.id},missing").get_json()
    assert body["active"] is True
    assert body["runs"][run.id]["label"] == "Fetching"
    assert f'id="job-{run.id}"' in body["runs"][run.id]["html"]


# ---- versions, preview, fact check ----

def test_new_version_keeps_earlier_files(app):
    client = app.test_client()
    run = _ready_job(app, client)
    assert [(d.kind, d.version) for d in app.store.list_documents(run.id)] == [("cv", 1)]

    page = client.get(f"/runs/{run.id}/questions").data
    assert b"New version" in page and b"Write v2" in page
    assert b'value="skip" checked' in page  # previous answers are pre-filled

    res = client.post(f"/runs/{run.id}/questions", data={"answer_0": "intermediate", "generate_letter": "1"})
    assert res.headers["Location"].endswith(f"/runs/{run.id}")
    docs = app.store.list_documents(run.id)
    assert sorted((d.kind, d.version) for d in docs) == [("cv", 1), ("cv", 2), ("letter", 2)]
    assert len({d.path for d in docs}) == 3 and all(Path(d.path).exists() for d in docs)

    with client.get(f"/download/{run.id}/cv") as latest:
        assert "CV_v2_" in latest.headers["Content-Disposition"]
    v1 = next(d for d in docs if d.version == 1)
    with client.get(f"/documents/{v1.id}/download") as old:
        assert old.status_code == 200 and "CV_v1_" in old.headers["Content-Disposition"]

    detail = client.get(f"/runs/{run.id}").data.decode()
    assert detail.count('class="dock-version"') == 2 and "v2 · latest" in detail  # every version, on the right
    assert "Confirmed for this version: Kubernetes" in detail
    strong = re.findall(r'<li class="m-strong">\s*<details[^>]*>\s*<summary>\s*<span class="m-term">([^<]+)<', detail)
    assert "Kubernetes" in strong  # the evidence reflects the level you gave


def test_preview_and_fact_check_are_shown(app):
    client = app.test_client()
    run = _ready_job(app, client, generate_letter="1")
    detail = client.get(f"/runs/{run.id}").data.decode()
    # Preview of exactly what was written, with the original wording available.
    assert 'data-version="1"' in detail and 'id="dock-body"' in detail and 'class="paper"' in detail  # on the right
    assert "Built Python APIs for Acme-style platforms" in detail and "Original: Built Python APIs" in detail
    assert 'data-tab="letter"' in detail and 'class="paper letter"' in detail and "I build Python APIs." in detail
    assert 'data-tab="ats"' in detail and "ATS match" in detail
    # Fact check explains the unconfirmed Kubernetes claims it removed.
    assert "Fact check" in detail
    assert "Kubernetes guru." in detail and "which your profile does not confirm" in detail
    assert "not a skill in your profile" in detail


def test_delete_one_version_moves_latest_back(app):
    client = app.test_client()
    run = _ready_job(app, client)
    client.post(f"/runs/{run.id}/questions", data={"answer_0": "skip"})
    v1, v2 = sorted(app.store.list_documents(run.id), key=lambda d: d.version)
    client.post(f"/documents/{v2.id}/delete")
    after = app.store.get_run(run.id)
    assert not Path(v2.path).exists() and after.output_cv_path == v1.path and after.status == "ready"
    client.post(f"/documents/{v1.id}/delete")
    after = app.store.get_run(run.id)
    assert after.output_cv_path == "" and after.status == "needs_answers"  # can write a new one


def test_delete_job_removes_its_files(app):
    client = app.test_client()
    run = _ready_job(app, client)
    path = Path(run.output_cv_path)
    res = client.post(f"/runs/{run.id}/delete")
    assert res.headers["Location"].endswith("/jobs")
    assert app.store.find_run(run.id) is None and not path.exists()


# ---- documents ----

def test_documents_lists_categorized_documents_with_filters(app):
    client = app.test_client()
    run = _ready_job(app, client)
    client.post(f"/runs/{run.id}/questions", data={"answer_0": "skip", "generate_letter": "1"})

    page = client.get("/documents").data.decode()
    assert 'Tailored CVs<span class="count">2</span>' in page
    assert 'Cover letters<span class="count">1</span>' in page
    assert 'Source CVs<span class="count">1</span>' in page
    assert client.get("/documents?kind=letter").data.count(b"/documents/") >= 1
    assert client.get("/documents?kind=cv&latest=1").data.count(b"older version") == 0
    assert b"No documents match" in client.get("/documents?q=nothing-matches").data
    assert b"Source CVs" in client.get("/documents?group=job").data
    assert client.get("/documents?company=Acme").data.count(b'/download"') == 3


def test_archiving_hides_documents_unless_requested(app):
    client = app.test_client()
    run = _ready_job(app, client)
    client.post(f"/runs/{run.id}/archive", data={"archived": "1"})
    assert b"No documents match" in client.get("/documents?kind=cv").data
    assert client.get("/documents?kind=cv&archived=1").data.count(b'/download"') == 1


def test_next_parameter_cannot_redirect_offsite(app):
    res = app.test_client().post("/jobs/bulk", data={"run_id": [], "next": "//evil.example.com/x"})
    assert res.headers["Location"].endswith("/jobs")


# ---- settings & your data ----

def test_settings_page_save_and_load(app):
    client = app.test_client()
    assert b"LLM_PROVIDER" in client.get("/settings").data
    res = client.post(
        "/settings",
        data={"LLM_PROVIDER": "groq", "LLM_MODEL": "llama-test", "LLM_API_KEY": "fake-key-1234",
              "AZURE_OPENAI_ENDPOINT": "", "AWS_DEFAULT_REGION": "", "OLLAMA_HOST": ""},
        follow_redirects=True,
    )
    assert b"Settings saved" in res.data
    res = client.get("/settings")
    assert b'value="groq" selected' in res.data and b"llama-test" in res.data
    assert b"fake-key-1234" not in res.data  # key is never echoed back
    assert b"1234" in res.data and b'type="password"' in res.data


def test_saved_settings_override_env_and_blank_key_keeps_saved(app, tmp_path, monkeypatch):
    from cv_maker.settings import effective_settings

    monkeypatch.setenv("LLM_PROVIDER", "openai")
    monkeypatch.setenv("LLM_API_KEY", "env-key")
    client = app.test_client()
    client.post("/settings", data={"LLM_PROVIDER": "openrouter", "LLM_MODEL": "m", "LLM_API_KEY": "saved-key"})
    client.post("/settings", data={"LLM_PROVIDER": "openrouter", "LLM_MODEL": "m2", "LLM_API_KEY": ""})
    values, sources = effective_settings(tmp_path)
    assert values["LLM_PROVIDER"] == "openrouter" and sources["LLM_PROVIDER"] == "settings"
    assert values["LLM_API_KEY"] == "saved-key" and values["LLM_MODEL"] == "m2"


def test_settings_test_endpoint_reports_problems(app, monkeypatch):
    monkeypatch.delenv("LLM_PROVIDER", raising=False)
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    body = app.test_client().post("/settings/test", data={"LLM_PROVIDER": "openrouter", "LLM_MODEL": "m", "LLM_API_KEY": ""}).get_json()
    assert body["ok"] is False and "No LLM configured" in body["message"]


def test_export_contains_everything_but_the_api_key(app):
    client = app.test_client()
    client.post("/settings", data={"LLM_PROVIDER": "openrouter", "LLM_MODEL": "m", "LLM_API_KEY": "sk-secret-123"})
    _ready_job(app, client, generate_letter="1")
    res = client.get("/settings/export")
    assert res.mimetype == "application/zip" and "cv-tailor-export-" in res.headers["Content-Disposition"]
    z = zipfile.ZipFile(io.BytesIO(res.data))
    names = z.namelist()
    assert {"README.txt", "profile.json", "jobs.json"} <= set(names)
    assert any(n.startswith("documents/") and n.endswith("v1-cv.docx") for n in names)
    assert any(n.startswith("documents/") and n.endswith("v1-cover-letter.docx") for n in names)
    assert any(n.startswith("source-cvs/") for n in names)
    assert json.loads(z.read("profile.json"))["name"] == "Ada Lovelace"
    assert all(b"sk-secret-123" not in z.read(n) for n in names)


def test_delete_all_requires_confirmation_and_removes_everything(app, tmp_path):
    client = app.test_client()
    client.post("/settings", data={"LLM_PROVIDER": "openrouter", "LLM_MODEL": "m", "LLM_API_KEY": "k"})
    _ready_job(app, client)
    client.post("/settings/delete-all", data={"confirm": "nope"})
    assert app.store.get_profile() is not None  # nothing happened

    client.post("/settings/delete-all", data={"confirm": "DELETE"})
    assert app.store.get_profile() is None and app.store.list_runs() == [] and app.store.list_uploads() == []
    assert list((tmp_path / "uploads").iterdir()) == [] and list((tmp_path / "output").iterdir()) == []
    assert (tmp_path / "settings.json").exists()  # kept unless asked

    client.post("/settings/delete-all", data={"confirm": "DELETE", "include_settings": "1"})
    assert not (tmp_path / "settings.json").exists()


def test_settings_page_explains_data_and_offers_control(app):
    page = app.test_client().get("/settings").data.decode()
    assert "Export my data" in page and "Delete all my data" in page and "Leaves this computer" in page


def test_delete_whole_version_removes_cv_and_letter(app):
    client = app.test_client()
    run = _ready_job(app, client, generate_letter="1")
    files = [Path(d.path) for d in app.store.list_documents(run.id)]
    assert len(files) == 2
    res = client.post(f"/runs/{run.id}/versions/1/delete")
    assert res.headers["Location"].endswith(f"/runs/{run.id}")
    assert app.store.list_documents(run.id) == [] and not any(f.exists() for f in files)
    after = app.store.get_run(run.id)
    assert (after.output_cv_path, after.letter_path, after.status) == ("", "", "needs_answers")
    assert client.post(f"/runs/{run.id}/versions/9/delete").status_code == 404
