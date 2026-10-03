"""Working in parallel without getting anyone banned: pacing, reuse, the read-only application check,
applying alongside the writing, and interview prep."""
import importlib.util
import json
import threading
import time

import httpx
import pytest

from cv_maker import prep
from cv_maker.apply import details as app_details
from cv_maker.jobs.fetch import fetch_job_url
from cv_maker.jobs.html import parse_job_page
from cv_maker.jobs.polite import SiteThrottle
from cv_maker.jobs.requirements import JobRequirements
from cv_maker.llm.chat import LLMError
from cv_maker.llm.limits import Governor
from cv_maker.models import Experience, Profile, Skill
from cv_maker.privacy import PrivateModel
from cv_maker.settings import llm_limits, save_automation, save_settings, suggested_limits
from cv_maker.store import Store


class Clock:
    def __init__(self) -> None:
        self.now = 1000.0
        self.slept = []

    def sleep(self, seconds: float) -> None:
        self.slept.append(round(seconds, 1))
        self.now += seconds


# ---- the LLM governor ----

def test_requests_stay_under_the_per_minute_budget():
    clock = Clock()
    governor = Governor(lambda: {"rpm": 2, "concurrency": 1, "daily": 0}, sleep=clock.sleep, clock=lambda: clock.now)
    for _ in range(3):
        governor.run(lambda: "ok")
    assert sum(clock.slept) >= 59  # the third waited for the first to fall out of the minute


def test_rate_limits_are_waited_out_as_asked_then_given_up_politely():
    clock, calls = Clock(), []

    def busy_once():
        calls.append(1)
        if len(calls) == 1:
            raise LLMError("OpenRouter rate-limited the request (HTTP 429).", status=429, retry_after=7)
        return "ok"

    governor = Governor(lambda: {"rpm": 0, "concurrency": 1, "daily": 0}, sleep=clock.sleep, clock=lambda: clock.now)
    assert governor.run(busy_once) == "ok" and clock.slept == [7]

    def always_busy():
        raise LLMError("busy (HTTP 503)", status=503)

    with pytest.raises(LLMError, match="stopped retrying"):
        governor.run(always_busy)
    assert len(clock.slept) == 1 + 3  # three waits, then it stops

    def bad_key():
        raise LLMError("rejected the API key (HTTP 401)", status=401)

    with pytest.raises(LLMError, match="401"):
        governor.run(bad_key)
    assert len(clock.slept) == 4  # never retried


def test_the_daily_cap_is_kept_across_calls():
    day = {}
    counter = type("C", (), {"get": lambda self: day.get("v"), "set": lambda self, v: day.update(v=v)})()
    governor = Governor(lambda: {"rpm": 0, "concurrency": 1, "daily": 2}, counter=counter)
    governor.run(lambda: 1)
    governor.run(lambda: 2)
    with pytest.raises(LLMError, match="today's limit of 2"):
        governor.run(lambda: 3)


def test_retries_count_toward_the_limits_and_used_up_quotas_are_not_retried():
    clock, day = Clock(), {}
    counter = type("C", (), {"get": lambda self: day.get("v"), "set": lambda self, v: day.update(v=v)})()
    governor = Governor(lambda: {"rpm": 0, "concurrency": 1, "daily": 5}, counter=counter, sleep=clock.sleep,
                        clock=lambda: clock.now)
    calls = []

    def busy_once():
        calls.append(1)
        if len(calls) == 1:
            raise LLMError("busy (HTTP 429)", status=429, retry_after=3)
        return "ok"

    assert governor.run(busy_once) == "ok" and day["v"]["count"] == 2  # the retry was a request too

    def quota_used_up():
        raise LLMError("rate-limited (HTTP 429)", status=429, retry_after=3600)

    with pytest.raises(LLMError, match="isn't retrying"):
        governor.run(quota_used_up)
    assert clock.slept == [3] and day["v"]["count"] == 3  # asked to wait an hour: no retries


def test_one_request_at_a_time_when_asked():
    governor = Governor(lambda: {"rpm": 0, "concurrency": 1, "daily": 0})
    inside, most = [0], [0]
    lock = threading.Lock()

    def call():
        with lock:
            inside[0] += 1
            most[0] = max(most[0], inside[0])
        time.sleep(0.05)
        with lock:
            inside[0] -= 1

    threads = [threading.Thread(target=governor.run, args=(call,)) for _ in range(4)]
    [t.start() for t in threads]
    [t.join() for t in threads]
    assert most[0] == 1


def test_suggested_limits_respect_free_tiers(tmp_path):
    assert suggested_limits("openrouter", "meta-llama/llama-3.3-70b-instruct:free")["daily"] == 45
    save_settings(tmp_path, {"LLM_PROVIDER": "openrouter", "LLM_MODEL": "x/y:free", "LLM_API_KEY": "k"})
    assert llm_limits(tmp_path)["rpm"] == 15
    save_automation(tmp_path, {"LLM_RPM": "5"})
    assert llm_limits(tmp_path)["rpm"] == 5 and llm_limits(tmp_path)["concurrency"] == 1


# ---- reuse of identical read-only answers ----

def test_identical_read_only_requests_reuse_the_earlier_answer_without_personal_data(tmp_path):
    store = Store(tmp_path / "db.sqlite")
    cache = type("Cache", (), {"get": lambda s, k: store.llm_cache_get(k), "put": lambda s, k, p, r: store.llm_cache_put(k, p, r)})()

    class Model:
        calls = 0

        def complete(self, messages, *, json_mode=False):
            Model.calls += 1
            return '{"job_title": "Engineer", "company": "Acme"}'

    log = []
    model = PrivateModel(Model(), list, lambda *e: log.append(e), cache=cache, scope="test")
    prompt = [{"role": "user", "content": "Read this job description and extract its requirements. Mail ada@example.com"}]
    assert json.loads(model.complete(prompt))["company"] == "Acme"
    assert json.loads(model.complete(prompt))["company"] == "Acme"
    assert Model.calls == 1 and log[1][3]["cached"] is True
    other = [{"role": "user", "content": "You are tailoring a CV to one job."}]
    model.complete(other)
    model.complete(other)
    assert Model.calls == 3  # writing is never reused
    import sqlite3
    from contextlib import closing
    with closing(sqlite3.connect(tmp_path / "db.sqlite")) as conn:
        assert "ada@example.com" not in str(conn.execute("SELECT * FROM llm_cache").fetchall())


# ---- job sites ----

def test_a_site_that_says_slow_down_is_left_alone():
    hits = []

    def handler(request):
        hits.append(str(request.url))
        return httpx.Response(999)

    throttle = SiteThrottle(intervals={}, default=0)
    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        first = fetch_job_url("https://www.linkedin.com/jobs/view/123456789/", client=client, throttle=throttle)
        second = fetch_job_url("https://www.linkedin.com/jobs/view/987654321/", client=client, throttle=throttle)
    assert "rate-limiting" in first.reason and len(hits) == 1  # no guest-API retry once the site refused
    assert "slow down" in second.reason and len(hits) == 1


def test_requests_to_one_site_are_spaced_out():
    clock = Clock()
    throttle = SiteThrottle(intervals={"linkedin.com": 6.0}, default=2.0, sleep=clock.sleep)
    throttle.wait("https://www.linkedin.com/jobs/view/1")
    throttle.wait("https://www.linkedin.com/jobs/view/2")
    throttle.wait("https://careers.acme.com/1")  # another site doesn't wait for LinkedIn
    assert len(clock.slept) == 1 and 5 <= clock.slept[0] <= 6


def test_requests_to_one_site_never_overlap():
    throttle = SiteThrottle(intervals={}, default=0)
    lock, inside, most = threading.Lock(), [0], [0]

    def request(url):
        with throttle.turn(url):
            with lock:
                inside[0] += 1
                most[0] = max(most[0], inside[0])
            time.sleep(0.05)  # a slow response
            with lock:
                inside[0] -= 1

    threads = [threading.Thread(target=request, args=(f"https://careers.acme.com/job/{i}",)) for i in range(4)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert most[0] == 1


def test_the_company_application_link_and_notes_come_from_the_posting():
    page = parse_job_page('<div class="show-more-less-html__markup">Build things</div><code id="applyUrl" style="display: none">'
                          '<!--"https://www.linkedin.com/jobs/view/externalApply/1?url=https%3A%2F%2Fcareers.acme.com%2Fjob%2F1'
                          '&amp;urlHash=x"--></code>')
    assert page.apply_url == "https://careers.acme.com/job/1"
    reqs = JobRequirements.from_dict({"notes": ["Applications close 15 October", "", "No visa sponsorship"] + ["x"] * 9})
    assert reqs.notes[:2] == ["Applications close 15 October", "No visa sponsorship"] and len(reqs.notes) == 8


# ---- interview prep ----

def _run(store, **kw):
    run = store.create_run(job_url="https://careers.acme.com/1", status="ready")
    run.title, run.company, run.jd_text = "Platform Engineer", "Acme", "Python and Kubernetes. Hybrid in London."
    run.requirements = {"must_have": [{"term": "Python", "original": "Python"}], "nice_to_have": [{"term": "Rust", "original": "Rust"}],
                        "notes": ["Hybrid, three days in London"]}
    run.gaps = [{"term": "Python", "status": "covered"}, {"term": "Kubernetes", "status": "missing"}]
    for key, value in kw.items():
        setattr(run, key, value)
    store.update_run(run)
    return run


def test_prep_notes_use_what_the_app_already_knows_in_one_request(tmp_path):
    store = Store(tmp_path / "db.sqlite")
    run = _run(store, scan={"questions": [{"label": "Notice period", "required": True, "saved": False}],
                            "documents": [{"what": "CV", "required": True}]})
    seen = []

    class Model:
        def complete(self, messages, *, json_mode=False):
            seen.append(messages[-1]["content"])
            return json.dumps({"role_in_brief": "Run the platform.", "brush_up": [{"topic": "Kubernetes", "why": "must-have", "status": "gap"}] * 12,
                               "likely_questions": [{"question": "How do you scale APIs?", "kind": "technical", "draw_on": "Northwind APIs"}],
                               "ask_them": ["What does on-call look like?"], "watch_outs": [], "company_points": ["Invented"]})

    notes = prep.write_notes(run, '{"experiences": []}', Model(), None)
    assert len(seen) == 1
    assert "Kubernetes (not in the profile)" in seen[0] and "Notice period" in seen[0] and "Hybrid, three days" in seen[0]
    assert len(notes["brush_up"]) == 8 and notes["company_points"] == []  # no company facts without a source


def _ada(store):
    store.save_profile(Profile(name="Ada Lovelace", email="ada@example.com", phone="", location="London", links={},
                               target_role=None, experiences=[Experience("Northwind", "Engineer", "London", "2019", None, True,
                                                                         ["Built Python APIs"], True)],
                               education=[], skills=[Skill("Python", "cv")], extras={}))


def test_prep_respects_the_research_switch_and_survives_failed_research(tmp_path):
    from cv_maker.tasks import Runner, Tasks

    store = Store(tmp_path / "db.sqlite")
    _ada(store)
    run = _run(store)
    asked = []

    def researcher(company, depth, job_url, force):
        asked.append(company)
        raise httpx.ConnectError("offline")

    class Model:
        def complete(self, messages, *, json_mode=False):
            return json.dumps({"role_in_brief": "Run the platform.", "brush_up": [], "likely_questions": [], "ask_them": [],
                               "watch_outs": [], "company_points": []})

    settings = {"RESEARCH_DEPTH": "thorough", "RESEARCH_AUTO": "0"}
    tasks = Tasks(store, None, lambda: Model(), None, Runner(sync=True), automation=lambda: settings, researcher=researcher)
    assert tasks.queue_prep(run)
    assert asked == [] and store.get_run(run.id).prep["notes"]["role_in_brief"] == "Run the platform."  # research stays off
    settings["RESEARCH_AUTO"] = "1"
    tasks.queue_prep(run)
    assert asked == ["Acme"] and store.get_run(run.id).prep["error"] == ""  # failed research didn't stop the notes


def test_long_pasted_reviews_say_how_much_was_summarised():
    seen = []

    class Model:
        def complete(self, messages, *, json_mode=False):
            seen.append(messages[-1]["content"])
            return json.dumps({"overall": "Mixed.", "pros": [], "cons": [], "interview_experiences": [], "ask_about": []})

    text = "Good team, slow releases. " * 700  # about 18,000 characters
    summary = prep.summarise_reviews(text, "Acme", Model())
    assert summary["chars"] == prep.REVIEWS_LIMIT and summary["pasted"] == len(text)
    assert text[:prep.REVIEWS_LIMIT] in seen[0] and text not in seen[0]


def test_background_work_shows_on_the_job_page_and_never_runs_twice(tmp_path):
    from cv_maker.app import create_app

    started, release, calls = threading.Event(), threading.Event(), []

    def researcher(company, depth, job_url, force):
        calls.append(depth)
        started.set()
        release.wait(10)

    app = create_app(data_dir=tmp_path, fetcher=lambda url: None, researcher=researcher)
    run = _run(app.store)
    client = app.test_client()
    client.post(f"/runs/{run.id}/research", data={"depth": "simple"})
    assert started.wait(5)
    page = client.get(f"/runs/{run.id}").data.decode()
    assert "Researching Acme" in page and 'name="depth"' not in page and 'data-busy="research"' in page
    assert client.get(f"/api/runs?ids={run.id}").get_json()["runs"][run.id]["busy"] == ["research"]
    client.post(f"/runs/{run.id}/research", data={"depth": "thorough"})
    assert calls == ["simple"]  # the second click didn't start another run
    release.set()
    end = time.monotonic() + 5
    while app.tasks.busy(run.id) and time.monotonic() < end:
        time.sleep(0.05)
    assert app.tasks.busy(run.id) == [] and 'name="depth"' in client.get(f"/runs/{run.id}").data.decode()


# ---- company research ----

def _web(robots: str = "User-agent: *\nAllow: /\n"):
    """A tiny internet: Wikipedia, Wikidata, an employer site and Hacker News. Records every URL asked for."""
    hits = []

    def handler(request):
        url = str(request.url)
        hits.append(url)
        host, path = request.url.host, request.url.path
        if host == "en.wikipedia.org" and path.endswith("api.php"):
            return httpx.Response(200, json=["Acme", ["Acme Robotics"], [""], [""]])
        if host == "en.wikipedia.org":
            return httpx.Response(200, json={"type": "standard", "title": "Acme Robotics", "description": "British robotics company",
                                             "extract": "Acme Robotics is a British company that builds warehouse robots.",
                                             "wikibase_item": "Q42", "content_urls": {"desktop": {"page": "https://en.wikipedia.org/wiki/Acme_Robotics"}}})
        if host == "www.wikidata.org" and "EntityData" in path:
            return httpx.Response(200, json={"entities": {"Q42": {"claims": {
                "P571": [{"mainsnak": {"datavalue": {"value": {"time": "+1999-01-01T00:00:00Z"}}}}],
                "P159": [{"mainsnak": {"datavalue": {"value": {"id": "Q84"}}}}],
                "P1128": [{"mainsnak": {"datavalue": {"value": {"amount": "+1200"}}}}],
                "P856": [{"mainsnak": {"datavalue": {"value": "https://www.acme-robotics.example"}}}]}}}})
        if host == "www.wikidata.org":
            return httpx.Response(200, json={"entities": {"Q84": {"labels": {"en": {"value": "London"}}}}})
        if host == "www.acme-robotics.example" and path == "/robots.txt":
            return httpx.Response(200, text=robots)
        if host == "www.acme-robotics.example" and path == "/":
            return httpx.Response(200, html='<a href="/about-us">About us</a><main>Home</main>')
        if host == "www.acme-robotics.example" and path == "/about-us":
            return httpx.Response(200, html="<main><h1>About</h1><p>We build robots that move parcels.</p></main>")
        if host == "hn.algolia.com":
            return httpx.Response(200, json={"hits": [
                {"title": "Acme Robotics raises Series C", "objectID": "1", "points": 120, "num_comments": 80, "created_at": "2026-05-01T00:00:00Z"},
                {"title": "Unrelated story", "objectID": "2", "points": 999, "num_comments": 1, "created_at": "2026-05-01T00:00:00Z"}]})
        return httpx.Response(404)

    return hits, httpx.Client(transport=httpx.MockTransport(handler), follow_redirects=True)


def test_simple_research_reads_only_wikipedia_and_links_the_rest(tmp_path):
    from cv_maker import research

    store = Store(tmp_path / "db.sqlite")
    hits, client = _web()
    with client:
        found = research.research("Acme Robotics", depth="simple", sources=["wikipedia", "wikidata", "glassdoor", "reddit"],
                                  store=store, client=client, throttle=SiteThrottle(intervals={}, default=0))
        again = research.research("acme  robotics", depth="simple", sources=["wikipedia"], store=store, client=client,
                                  throttle=SiteThrottle(intervals={}, default=0))
    assert found["basics"]["summary"].startswith("Acme Robotics is a British company") and found["facts"] == {}
    assert {link["site"] for link in found["links"]} == {"Glassdoor", "Reddit"}
    assert all("duckduckgo.com" in link["url"] for link in found["links"])
    assert not [h for h in hits if "glassdoor" in h or "reddit" in h]  # linked, never fetched
    assert again == found and len(hits) == 1  # one request (the company's own page), and nothing re-fetched later


def test_thorough_research_gathers_sourced_facts_and_respects_robots(tmp_path):
    from cv_maker import research

    store = Store(tmp_path / "db.sqlite")
    asked = []

    class Model:
        def complete(self, messages, *, json_mode=False):
            asked.append(messages[-1]["content"])
            return json.dumps({"what_they_do": "Warehouse robots.", "points": [{"point": "Founded in 1999", "source": "Wikidata"}],
                               "discussions": [{"point": "A large funding round", "source": "Hacker News"}], "for_interviews": ["Ask about scale"]})

    sources = ["wikipedia", "wikidata", "website", "hackernews"]
    hits, client = _web()
    research._robots.clear()
    with client:
        found = research.research("Acme Robotics", depth="thorough", sources=sources, store=store, get_model=lambda: Model(),
                                  client=client, throttle=SiteThrottle(intervals={}, default=0))
    assert found["facts"] == {"Founded": "1999", "Headquarters": "London", "Employees": "1,200", "Website": "https://www.acme-robotics.example"}
    assert "We build robots that move parcels." in found["website"]["text"]
    assert [d["title"] for d in found["discussions"]] == ["Acme Robotics raises Series C"]
    assert found["summary"]["points"] == [{"point": "Founded in 1999", "source": "Wikidata"}] and len(asked) == 1
    assert "Organise what was found about Acme Robotics" in asked[0]

    hits, client = _web(robots="User-agent: *\nDisallow: /\n")
    research._robots.clear()
    with client:
        blocked = research.research("Acme Robotics", depth="thorough", sources=sources, store=store, get_model=lambda: Model(),
                                    client=client, throttle=SiteThrottle(intervals={}, default=0), force=True)
    assert blocked["website"]["allowed"] is False
    assert not [h for h in hits if h.endswith("/about-us") or h == "https://www.acme-robotics.example/"]
    assert {"source": "Company website", "why": "its robots.txt asks automated visitors not to read it"} in blocked["skipped"]


def test_wikidata_facts_are_the_current_ones():
    from cv_maker import research

    def handler(request):
        if "EntityData" in request.url.path:
            return httpx.Response(200, json={"entities": {"Q1": {"claims": {
                "P169": [{"rank": "normal", "mainsnak": {"datavalue": {"value": {"id": "Q10"}}}, "qualifiers": {"P582": [{}]}},
                         {"rank": "normal", "mainsnak": {"datavalue": {"value": {"id": "Q11"}}}}],
                "P1128": [{"rank": "normal", "mainsnak": {"datavalue": {"value": {"amount": "+500"}}},
                           "qualifiers": {"P585": [{"datavalue": {"value": {"time": "+2010-01-01T00:00:00Z"}}}]}},
                          {"rank": "normal", "mainsnak": {"datavalue": {"value": {"amount": "+900"}}},
                           "qualifiers": {"P585": [{"datavalue": {"value": {"time": "+2024-01-01T00:00:00Z"}}}]}}]}}}})
        return httpx.Response(200, json={"entities": {"Q11": {"labels": {"en": {"value": "Current CEO"}}}}})

    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        facts = research.wikidata_facts("Q1", client, SiteThrottle(intervals={}, default=0))
    assert facts == {"Chief executive": "Current CEO", "Employees": "900 (2024)"}  # not the past CEO or the 2010 headcount


def test_a_common_name_finds_the_company_not_the_fruit():
    from cv_maker import research
    asked = []

    def handler(request):
        asked.append(request.url.path)
        if request.url.path.endswith("api.php"):
            return httpx.Response(200, json=["Apple", ["Apple", "Apple Inc.", "Apple Records"], ["", "", ""], ["", "", ""]])
        if request.url.path.endswith("/Apple"):
            return httpx.Response(200, json={"type": "standard", "title": "Apple", "description": "Fruit",
                                             "extract": "An apple is a round, edible fruit."})
        return httpx.Response(200, json={"type": "standard", "title": "Apple Inc.", "description": "American technology company",
                                         "extract": "Apple Inc. is an American multinational technology company.", "wikibase_item": "Q312"})

    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        found = research.wikipedia("Apple", client, SiteThrottle(intervals={}, default=0))
    assert found["title"] == "Apple Inc." and found["wikidata"] == "Q312" and len(asked) == 3


def test_tavily_brings_news_and_what_people_say_with_your_own_key(tmp_path):
    from cv_maker import research

    store, asked, sent = Store(tmp_path / "db.sqlite"), [], []

    def handler(request):
        assert request.headers["Authorization"] == "Bearer tvly-test-0123456789"
        body = json.loads(request.content)
        sent.append(body)
        if body["topic"] == "news":
            return httpx.Response(200, json={"results": [{"title": "Acme opens a Leeds office", "url": "https://www.reuters.com/acme",
                                                          "content": "Acme Robotics will hire 200 engineers.", "published_date": "2026-09-01"}]})
        return httpx.Response(200, json={"results": [{"title": "Acme Robotics reviews", "url": "https://www.glassdoor.com/acme",
                                                      "content": "Good mentoring, slow promotions."}]})

    class Model:
        def complete(self, messages, *, json_mode=False):
            asked.append(messages[-1]["content"])
            return json.dumps({"what_they_do": "Robots.", "points": [{"point": "Opening a Leeds office", "source": "reuters.com"}],
                               "what_people_say": [{"point": "Good mentoring", "source": "glassdoor.com via Tavily"}]})

    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        found = research.research("Acme Robotics", depth="thorough", sources=["tavily"], store=store, get_model=lambda: Model(),
                                  client=client, throttle=SiteThrottle(intervals={}, default=0), tavily_key="tvly-test-0123456789")
    assert [b["query"] for b in sent] == ["Acme Robotics", "Acme Robotics employee reviews, work culture and interview process"]
    assert all("Ada" not in json.dumps(b) for b in sent)  # only the company's name goes to Tavily
    assert found["web"]["news"][0]["site"] == "reuters.com" and found["web"]["people"][0]["site"] == "glassdoor.com"
    assert "Tavily" in found["used"] and "Acme opens a Leeds office" in asked[0] and "Good mentoring, slow promotions." in asked[0]
    assert found["summary"]["what_people_say"] == [{"point": "Good mentoring", "source": "glassdoor.com via Tavily"}]


def test_tavily_without_a_key_or_with_a_rejected_key_is_skipped_and_said(tmp_path):
    from cv_maker import research

    store = Store(tmp_path / "db.sqlite")
    with httpx.Client(transport=httpx.MockTransport(lambda r: httpx.Response(401))) as client:
        no_key = research.research("Acme", depth="thorough", sources=["tavily"], store=store, client=client,
                                   throttle=SiteThrottle(intervals={}, default=0))
        bad_key = research.research("Acme", depth="thorough", sources=["tavily"], store=store, client=client, force=True,
                                    throttle=SiteThrottle(intervals={}, default=0), tavily_key="tvly-wrong-0123456789")
    assert {"source": "Tavily", "why": "add your Tavily API key in Settings to use it"} in no_key["skipped"]
    assert any(s["source"] == "Tavily" and "rejected the API key" in s["why"] for s in bad_key["skipped"])
    assert no_key["web"] is None and bad_key["web"] is None


def test_your_tavily_key_is_checked_saved_and_never_shown(tmp_path, monkeypatch):
    from cv_maker import research
    from cv_maker.app import create_app
    from cv_maker.export.archive import build_export

    app = create_app(data_dir=tmp_path, fetcher=lambda url: None)
    client = app.test_client()
    page = client.post("/settings/automation", data={"TAVILY_API_KEY": "not-a-key", "sources_shown": "1"},
                       follow_redirects=True).data.decode()
    assert "doesn&#39;t look like a Tavily key" in page
    page = client.post("/settings/automation", data={"TAVILY_API_KEY": "tvly-dev-abcdef123456", "sources_shown": "1",
                                                     "RESEARCH_SOURCES": "wikipedia"}, follow_redirects=True).data.decode()
    assert "Tavily search is now one of the research sources" in page and "…3456" in page and "abcdef123456" not in page
    from cv_maker.settings import automation_settings, tavily_key
    assert tavily_key(tmp_path) == ("tvly-dev-abcdef123456", "settings")
    assert automation_settings(tmp_path)["RESEARCH_SOURCES"] == "wikipedia,tavily"
    monkeypatch.setattr(research, "tavily_usage", lambda key: (True, f"The key works: checked {key[-4:]}"))
    assert client.post("/settings/tavily/test", data={}).get_json() == {"ok": True, "message": "The key works: checked 3456"}
    import zipfile
    with zipfile.ZipFile(build_export(app.store)) as z:
        assert not any("tvly-dev" in z.read(name).decode("utf-8", "replace") for name in z.namelist())
    import logging
    logging.getLogger("cv_maker.test").warning("calling Tavily with tvly-dev-abcdef123456")
    assert "tvly-dev-abcdef123456" not in (tmp_path / "logs" / "cv-tailor.log").read_text(encoding="utf-8")
    client.post("/settings/automation", data={"clear_tavily_key": "1"})
    assert tavily_key(tmp_path) == ("", "")


def test_a_source_that_says_slow_down_is_left_alone(tmp_path):
    from cv_maker import research

    hits = []

    def handler(request):
        hits.append(str(request.url))
        return httpx.Response(429, text="You are making too many requests")

    store = Store(tmp_path / "db.sqlite")
    throttle = SiteThrottle(intervals={}, default=0)
    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        found = research.research("Acme Robotics", depth="simple", sources=["wikipedia"], store=store, client=client, throttle=throttle)
        again = research.research("Acme Robotics", depth="simple", sources=["wikipedia"], store=store, client=client,
                                  throttle=throttle, force=True)
    assert found["basics"] is None and "slow down (HTTP 429)" in found["skipped"][0]["why"]
    assert "left alone for now" in again["skipped"][0]["why"] and len(hits) == 1  # not asked a second time


def test_politeness_levels_set_every_pace():
    from cv_maker.jobs import polite

    polite.configure(lambda: "gentle")
    try:
        assert polite.level()["between"] == 4.0 and polite.level()["page_pause_ms"] == 2500
        clock = Clock()
        throttle = SiteThrottle(sleep=clock.sleep)
        throttle.wait("https://careers.acme.com/1")
        throttle.wait("https://careers.acme.com/2")
        assert 3.5 <= clock.slept[0] <= 4.0
    finally:
        polite.configure(lambda: "standard")


# ---- the read-only look and applying alongside (real browser, mock job site) ----

def _browser_available() -> bool:
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as pw:
            pw.chromium.launch(headless=True).close()
        return True
    except Exception:
        return False


needs_browser = pytest.mark.skipif(not _browser_available(), reason="Playwright with Chromium is not installed")
DETAILS = {"standard": {"first_name": "Ada", "last_name": "Lovelace", "email": "ada@example.com", "country": "United Kingdom",
                        "work_authorization": "Yes", "sponsorship": "No", "salary": "£70,000"}, "custom": []}


@needs_browser
def test_the_look_reports_what_an_application_needs_without_entering_anything():
    from fixtures.mock_ats import Running

    from cv_maker.apply import scout

    site = Running()
    try:
        walled = scout.look(f"{site.url}/job/1", DETAILS, throttle=SiteThrottle(intervals={}, default=0))
        form = scout.look(f"{site.url}/apply/1/experience", DETAILS, throttle=SiteThrottle(intervals={}, default=0))
    finally:
        site.close()
    assert walled["stopped_at"] == "account" and walled["account"]["options"] == ["email"]
    assert walled["account"]["creating"] is False and site.app.received == {}  # nothing was entered anywhere
    assert form["stopped_at"] == "form"
    assert [(d["what"], d["required"]) for d in form["documents"]] == [("CV", True), ("Cover letter", False)]
    questions = {q["label"]: q["saved"] for q in form["questions"]}
    assert questions["Years of experience with Python *"] is False  # the user will be asked
    assert "needs an account on 127.0.0.1" in " ".join(scout.summary(walled)) or "needs an account" in " ".join(scout.summary(walled))


@needs_browser
def test_applying_alongside_fills_what_it_can_waits_for_the_cv_and_asks_for_a_required_letter(tmp_path):
    from fixtures.mock_ats import Running

    from cv_maker.apply.session import ApplySession

    site = Running()
    store = Store(tmp_path / "db.sqlite")
    store.save_profile(Profile(name="Ada Lovelace", email="ada@example.com", phone="", location="London", links={},
                               target_role=None, experiences=[Experience("Northwind", "Engineer", "", "2019", None, True, ["Built APIs"], True)],
                               education=[], skills=[Skill("Python", "cv")], extras={}))
    app_details.save(store, DETAILS["standard"])
    run = store.create_run(job_url=f"{site.url}/apply/3/info", status="generating")
    application = store.create_application(run, 1, f"{site.url}/apply/3/info")
    ready = threading.Event()
    cv = tmp_path / "cv_v1.docx"
    cv.write_bytes(b"cv")

    class Model:
        def complete(self, messages, *, json_mode=False):
            prompt = messages[-1]["content"]
            import re
            field = re.search(r"^(\S+) · \w+ · Years of experience with Python", prompt, re.M)
            button = re.search(r"^(\S+) · Save and Continue$", prompt, re.M)
            return json.dumps({"fill": [{"id": field.group(1), "value": "6"}] if field else [],
                               "next": {"do": "click", "id": button.group(1)} if button else {"do": "wait"}})

    session = ApplySession(application=application, store=store, get_model=lambda: PrivateModel(Model()),
                           browser_dir=tmp_path / "browser", files=lambda: {"cv": str(cv), "letter": ""} if ready.is_set() else None,
                           job_label="Platform Engineer", job_context="Python")
    try:
        session.start()
        end = time.monotonic() + 60
        while time.monotonic() < end and (store.get_application(application.id).waiting or {}).get("kind") != "documents":
            time.sleep(0.2)
        waiting = store.get_application(application.id)
        assert waiting.waiting["kind"] == "documents" and waiting.status == "working"  # not "needs you": it's waiting
        assert site.app.received["info"]["first"] == "Ada"  # the first page was done while the CV was being written
        assert store.find_run(run.id).letter_required is True  # the writer was told the form needs a letter
        ready.set()  # the CV is written, but this version has no letter
        while time.monotonic() < end and (store.get_application(application.id).waiting or {}).get("kind") != "letter":
            time.sleep(0.2)
        assert store.get_application(application.id).waiting["kind"] == "letter"
    finally:
        session.send({"action": "stop"})
        session.join(timeout=30)
        site.close()


# ---- in the app ----

def test_the_app_checks_applications_after_analysis_and_writes_prep_after_the_cv(tmp_path):
    from io import BytesIO

    from docx import Document

    from cv_maker.app import create_app
    from cv_maker.jobs.fetch import FetchResult

    looked, alongside = [], []

    class Model:
        def complete(self, messages, *, json_mode=False):
            prompt = messages[-1]["content"]
            if prompt.startswith("Extract a structured profile"):
                return json.dumps({"name": "Ada Lovelace", "email": "ada@example.com", "experiences": [
                    {"company": "Northwind", "title": "Engineer", "bullets": ["Built Python APIs"]}], "skills": ["Python"]})
            if prompt.startswith("Read this job description"):
                return json.dumps({"job_title": "Platform Engineer", "company": "Acme", "must_have": [{"term": "Python", "original": "Python"}],
                                   "notes": ["Applications close 15 October"]})
            if prompt.startswith("You are tailoring"):
                return json.dumps({"summary": "Python engineer.", "skills": ["Python"],
                                   "experiences": [{"company": "Northwind", "title": "Engineer", "bullets": ["Built Python APIs"]}]})
            if prompt.startswith("You are helping a candidate prepare"):
                return json.dumps({"role_in_brief": "Keep the platform running.", "brush_up": [{"topic": "Kubernetes", "why": "listed", "status": "gap"}],
                                   "likely_questions": [{"question": "Tell me about an API you built.", "kind": "behavioural", "draw_on": "Northwind"}],
                                   "ask_them": ["How big is the team?"], "watch_outs": ["Closing date"], "company_points": []})
            if prompt.startswith("Summarise these reviews"):
                return json.dumps({"overall": "Good teams, slow processes.", "pros": ["Good teams"], "cons": ["Slow"],
                                   "interview_experiences": ["Two rounds"], "ask_about": ["Release cadence"]})
            return "{}"

    def scanner(url, details):
        looked.append(url)
        return {"checked_at": "2026-10-03T10:00:00+00:00", "start_url": url, "sites": ["acme.com", "myworkdayjobs.com"],
                "account": {"site": "myworkdayjobs.com", "options": ["email", "Google"], "creating": False}, "captcha": False,
                "documents": [{"what": "Cover letter", "label": "Cover letter", "required": True}],
                "questions": [{"label": "Notice period", "required": True, "saved": False}], "consent": [], "diversity": True,
                "stopped_at": "account", "links": [{"text": "Glassdoor reviews", "href": "https://www.glassdoor.com/acme"}]}

    researched = []

    def researcher(company, depth, job_url, force):
        researched.append((company, depth))
        found = {"company": company, "depth": depth, "researched_at": "2026-10-03T10:00:00+00:00", "facts": {"Founded": "1999"},
                 "basics": {"summary": "Acme makes anvils.", "url": "https://en.wikipedia.org/wiki/Acme", "title": "Acme"},
                 "summary": None, "discussions": [], "website": None, "used": ["Wikipedia"], "skipped": [], "links": []}
        app.store.save_research("acme", found)
        return found

    fetch = lambda url: FetchResult(ok=True, text="Need Python. Apply by 15 October.", reason="", title="Platform Engineer", company="Acme")  # noqa: E731
    app = create_app(data_dir=tmp_path, model_factory=Model, fetcher=fetch, sync_jobs=True, scanner=scanner, researcher=researcher)
    app.tasks.on_generate = lambda run, version: alongside.append((run.id, version))
    client = app.test_client()
    buf = BytesIO()
    doc = Document()
    doc.add_paragraph("Ada Lovelace, Python engineer")
    doc.add_paragraph("Engineer at Northwind since 2021. Built Python APIs used by three teams and ran on-call for the payments platform.")
    doc.save(buf)
    res = client.post("/profile/upload", data={"cv": (BytesIO(buf.getvalue()), "cv.docx")}, content_type="multipart/form-data")
    client.post(res.headers["Location"], data={"action": "send"})
    client.post("/jobs", data={"urls": "https://careers.acme.com/job/1"})
    run = app.store.list_runs()[0]
    assert looked == ["https://careers.acme.com/job/1"] and app.store.get_run(run.id).scan["account"]["site"] == "myworkdayjobs.com"
    assert researched == [("Acme", "simple")]  # the research agent ran by itself, at the chosen depth

    questions = client.get(f"/runs/{run.id}/questions").data.decode()
    if importlib.util.find_spec("playwright"):
        assert 'name="apply_alongside"' in questions and "Fill in the application on acme.com while" in questions
    else:  # the optional browser engine isn't installed: the toggle is replaced by the reason
        assert 'name="apply_alongside"' not in questions and "the browser engine isn&#39;t installed" in questions
    client.post(f"/runs/{run.id}/questions", data={"apply_alongside": "1"})
    assert alongside == [(run.id, 1)]
    done = app.store.get_run(run.id)
    assert done.status == "ready" and done.prep["notes"]["brush_up"][0]["topic"] == "Kubernetes"

    page = client.get(f"/runs/{run.id}").data.decode()
    assert "<h2>Applying</h2>" in page and "Applications close 15 October" in page
    assert "Needs you to sign in on myworkdayjobs.com (email, Google)" in page and "You'll be asked: Notice period" in page
    assert "none written yet" in page and "Glassdoor reviews ↗" in page
    assert "Interview prep" in page and "Tell me about an API you built." in page and "duckduckgo.com" in page
    assert "Company research: Acme" in page and "Acme makes anvils." in page and "<dt>Founded</dt>" in page
    client.post(f"/runs/{run.id}/research", data={"depth": "thorough"})
    assert researched[-1] == ("Acme", "thorough")

    client.post(f"/runs/{run.id}/prep/reviews", data={"reviews": "Great people and interesting work, but releases are slow. " * 3})
    assert "Good teams, slow processes." in client.get(f"/runs/{run.id}").data.decode()

    client.post("/settings/automation", data={"LLM_RPM": "10", "PREP_AUTO": "1", "RESEARCH_DEPTH": "thorough", "POLITENESS": "gentle",
                                              "sources_shown": "1", "RESEARCH_SOURCES": ["wikipedia", "ambitionbox", "nonsense"]})
    from cv_maker.settings import automation_settings
    saved = automation_settings(tmp_path)
    assert (saved["RESEARCH_DEPTH"], saved["POLITENESS"], saved["RESEARCH_SOURCES"]) == ("thorough", "gentle", "wikipedia,ambitionbox")
    assert saved["SCAN_APPLICATIONS"] == "0"  # unticked
    settings = client.get("/settings").data.decode()
    assert 'id="automation"' in settings and 'value="10"' in settings and "10 a minute" in settings
    assert 'id="polite-info"' in settings and "Time off after a site says slow down" in settings and "15 min" in settings
    assert "Linked only" in settings and 'value="ambitionbox" checked' in settings
