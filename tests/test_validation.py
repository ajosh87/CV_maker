"""Validation: what you type is checked before it's saved or sent, and every check says what to fix."""
import json
from io import BytesIO

import pytest
from docx import Document

from cv_maker.app import create_app
from cv_maker.export.docx_cv import ExportError, verify_docx
from cv_maker.jobs.fetch import FetchResult
from cv_maker.llm.chat import LLMError, complete_json
from cv_maker.models import Experience, Profile, Skill
from cv_maker.privacy import PrivateModel
from cv_maker.settings import automation_settings, load_saved_settings


@pytest.fixture
def app(tmp_path):
    fetch = lambda url: FetchResult(ok=True, text="Need Python. " * 20, reason="", title="Engineer", company="Acme")  # noqa: E731
    app = create_app(data_dir=tmp_path, fetcher=fetch, sync_jobs=True)
    app.store.save_profile(Profile(name="Ada Lovelace", email="ada@example.com", phone="", location="London", links={},
                                   target_role=None, experiences=[Experience("Northwind", "Engineer", "London", "2019", None,
                                                                             True, ["Built APIs"], True)],
                                   education=[], skills=[Skill("Python", "cv")], extras={}))
    return app


def test_llm_settings_say_what_is_wrong_and_keep_what_you_typed(app, tmp_path):
    client = app.test_client()
    res = client.post("/settings", data={"LLM_PROVIDER": "azure", "LLM_MODEL": "gpt 4o", "LLM_API_KEY": "sk-abc def",
                                         "AZURE_OPENAI_ENDPOINT": "my resource"})
    page = res.data.decode()
    assert res.status_code == 400 and "Nothing was saved" in page
    assert "The model name has spaces in it" in page and "Azure endpoint as a link" in page and "spaces or line breaks" in page
    assert 'value="gpt 4o"' in page and "sk-abc def" not in page  # what you typed stays, except the key
    assert not (tmp_path / "settings.json").exists()


def test_pace_numbers_must_be_whole_numbers_in_range(app, tmp_path):
    client = app.test_client()
    page = client.post("/settings/automation", data={"LLM_RPM": "lots", "LLM_CONCURRENCY": "0"}, follow_redirects=True).data.decode()
    assert "Requests a minute: enter a whole number" in page and "At the same time: enter a whole number from 1" in page
    assert automation_settings(tmp_path)["LLM_RPM"] == ""  # nothing saved


def test_words_too_short_to_hide_safely_are_refused(app):
    client = app.test_client()
    page = client.post("/settings/privacy", data={"never_send": "a\nAcme Secret Project\nacme secret project"},
                       follow_redirects=True).data.decode()
    assert "“a” is too short to hide" in page
    assert app.store.get_never_send() == ["Acme Secret Project"]


def test_application_details_are_checked_and_what_you_typed_is_kept(app):
    client = app.test_client()
    res = client.post("/settings/apply", data={"email": "ada at example", "phone": "call me", "linkedin": "linkedin/ada",
                                               "salary": "£70,000"})
    page = res.data.decode()
    assert res.status_code == 400 and "doesn&#39;t look like an email address" in page and "Phone: use digits" in page
    assert "LinkedIn URL: enter the full link" in page and 'value="£70,000"' in page
    assert app.store.get_preference("apply_details") is None


def test_profile_edits_need_a_name_and_real_links(app):
    client = app.test_client()
    res = client.post("/profile/edit", data={"name": "", "email": "ada@example.com", "links": "github: not a link at all"})
    page = res.data.decode()
    assert res.status_code == 400 and "Name: it goes at the top of every CV" in page and "enter a web address" in page
    assert app.store.get_profile().name == "Ada Lovelace"  # unchanged


def test_uploads_are_checked_for_size_and_for_real_text(app):
    client = app.test_client()
    big = client.post("/profile/upload", data={"cv": (BytesIO(b"0" * (11 * 1024 * 1024)), "cv.pdf")},
                      content_type="multipart/form-data", follow_redirects=True)
    assert "larger than 10 MB" in big.data.decode()
    buf = BytesIO()
    doc = Document()
    doc.add_paragraph("Ada")
    doc.save(buf)
    tiny = client.post("/profile/upload", data={"cv": (BytesIO(buf.getvalue()), "scan.docx")}, content_type="multipart/form-data")
    assert tiny.status_code == 400 and "Only 3 characters of text were found" in tiny.data.decode()


def test_a_large_batch_of_links_is_capped_and_said(app):
    client = app.test_client()
    urls = "\n".join(f"https://careers.acme.com/job/{i}" for i in range(30))
    page = client.post("/jobs", data={"urls": urls}, follow_redirects=True).data.decode()
    assert len(app.store.list_runs()) == 25 and "5 more were left out" in page


def test_a_broken_json_reply_gets_one_more_chance_and_is_never_reused():
    calls, stored = [], {}

    class Flaky:
        def complete(self, messages, *, json_mode=False):
            calls.append(messages)
            return "Sure! Here it is: {broken" if len(calls) % 2 else json.dumps({"job_title": "Engineer"})

    cache = type("Cache", (), {"get": lambda self, k: stored.get(k), "put": lambda self, k, purpose, reply: stored.update({k: reply})})()
    model = PrivateModel(Flaky(), lambda: [], cache=cache)
    assert complete_json(model, "Read this job description and list the requirements: Python") == {"job_title": "Engineer"}
    assert len(calls) == 2 and calls[1][-1]["content"].startswith("Your last reply was not valid JSON")
    assert list(stored.values()) == []  # the broken reply wasn't cached; the repaired one wasn't a read-only request

    class Hopeless:
        def complete(self, messages, *, json_mode=False):
            return "no json here"

    with pytest.raises(LLMError, match="not valid JSON"):
        complete_json(Hopeless(), "Read this job description: Python")


def test_a_docx_that_cant_be_opened_again_is_an_error_not_a_download(tmp_path):
    broken = tmp_path / "cv.docx"
    broken.write_bytes(b"PK\x03\x04 not really a document")
    with pytest.raises(ExportError, match="couldn't be opened again"):
        verify_docx(broken)
    good = tmp_path / "ok.docx"
    doc = Document()
    doc.add_paragraph("Ada Lovelace")
    doc.save(good)
    verify_docx(good, "Ada Lovelace")
    with pytest.raises(ExportError, match="missing its text"):
        verify_docx(good, "Someone Else")


def test_settings_are_written_in_one_step(tmp_path):
    from cv_maker.settings import save_settings

    save_settings(tmp_path, {"LLM_PROVIDER": "openrouter", "LLM_MODEL": "x/y:free", "LLM_API_KEY": "k-123"})
    assert load_saved_settings(tmp_path)["LLM_API_KEY"] == "k-123"
    assert [p.name for p in tmp_path.iterdir() if p.suffix == ".tmp"] == []  # no half-written files left behind
